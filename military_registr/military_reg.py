import os
from openpyxl import load_workbook

from name_utils import get_fio_part, get_name_patronymic_initials


def _get_relatives_list(parsed_data_1c: dict) -> list[dict]:
    personal_info = parsed_data_1c.get('personal_info', {})
    relatives_list = personal_info.get('РодственникиСписок') or []
    if relatives_list:
        return relatives_list[:5]

    single_relative = personal_info.get('Родственники', {})
    if any(single_relative.get(key, '') for key in ('Фамилия', 'Имя', 'Отчество', 'ДатаРождения', 'СтепеньРодства')):
        return [single_relative]
    return []


def _get_family_input_words(parsed_data_1c: dict) -> dict:
    relatives_list = _get_relatives_list(parsed_data_1c)
    placeholders = (
        ('СтепеньРодства', 'Степень родства'),
        ('Фамилия', 'Фамилия родственника'),
        ('Имя', 'Имя родственника'),
        ('Отчество', 'Отчество родственника'),
        ('ДатаРождения', 'Дата рождения родственника'),
    )
    family_input_words: dict = {}

    for index in range(5):
        relative = relatives_list[index] if index < len(relatives_list) else {}
        suffix = '' if index == 0 else str(index + 1)
        for source_key, placeholder_name in placeholders:
            family_input_words[f'<{placeholder_name}{suffix}>'] = relative.get(source_key, '')

    return family_input_words


def _shared_military_registration_input_words(parsed_data_1c: dict) -> dict:
    personal_info = parsed_data_1c.get('personal_info', {})
    job_info = parsed_data_1c.get('job_info', {})
    document_info = parsed_data_1c.get('document_info', {})
    military_info = parsed_data_1c.get('military_info', {})
    org_info = job_info.get('Организация', {})
    driver_license = personal_info.get('ВодительскоеУдостоверение', {})

    shared_input_words = {
        '<Фамилия>': get_fio_part(personal_info.get('ФИО', ''), 0),
        '<Имя>': get_fio_part(personal_info.get('ФИО', ''), 1),
        '<Отчество>': get_fio_part(personal_info.get('ФИО', ''), 2),
        '<Полное имя работника>': personal_info.get('ФИО', ''),
        '<Полное наименование>': org_info.get('НаименованиеПолное', ''),
        '<Краткое наименование>': org_info.get('НаименованиеСокращенное', ''),
        '<Юридический адрес>': org_info.get('ЮридическийАдрес', ''),
        '<Фактический адрес>': org_info.get('ФактическийАдрес', ''),
        '<ИНН организации>': org_info.get('ИНН_Орг', ''),
        '<КПП организации>': org_info.get('КПП_Орг', ''),
        '<ОГРН организации>': org_info.get('ОГРН(ИП)', ''),
        '<Дата рождения>': document_info.get('ДатаРождения', ''),
        '<Место рождения>': document_info.get('МестоРождения', ''),
        '<Семейное положение>': personal_info.get('СемейноеПоложение', ''),
        '<Адрес места проживания>': personal_info.get('Адрес_проживания', ''),
        '<Адрес по прописке>': personal_info.get('Адрес_прописка', ''),
        '<Мобильный телефон>': personal_info.get('Телефон', ''),
        '<СНИЛС>': personal_info.get('СНИЛС', ''),
        '<Подразделение>': job_info.get('Подразделение', ''),
        '<Должность работника>': job_info.get('Должность', ''),
        '<Дата приказа о приеме>': job_info.get('ДатаПриемаНаРаботу', ''),
        '<Номер приказа о приеме>': job_info.get('НомерПриказаОПриеме', job_info.get('ТкНомер', '')),
        '<Приказ №>': job_info.get('НомерПриказаОПриеме', job_info.get('ТкНомер', '')),
        '<Договор №>': job_info.get('ТкНомер', ''),
        '<Военкомат>': military_info.get('Военкомат', ''),
        '<Адрес военкомата>': military_info.get('АдресВоенкомата', ''),
        '<Категория запаса>': military_info.get('КатегорияЗапаса', ''),
        '<Воинское звание>': military_info.get('ВоинскоеЗвание', ''),
        '<Состав (профиль)>': military_info.get('СоставПрофиль', ''),
        '<ВУС>': military_info.get('ВУС', ''),
        '<Годность к воинской службе>': military_info.get('Годность', ''),
        '<Документ воинского учёта>': military_info.get('ДокументВоинскогоУчета', ''),
        '<Серия документа воинского учёта>': military_info.get('СерияДокументаВоинскогоУчета', ''),
        '<Номер документа воинского учёта>': military_info.get('НомерДокументаВоинскогоУчета', ''),
        '<Отдел, выдавший документ воинского учёта>': military_info.get('ОтделВыдавшийДокументВоинскогоУчета', ''),
        '<Дата выдачи документа воинского учёта>': military_info.get('ДатаВыдачиДокументаВоинскогоУчета', ''),
        '<Должность ответственного за ВУ>': military_info.get('ДолжностьОтветственногоЗаВУ', ''),
        '<Инициалы имени, отчества ответственного за ВУ>': military_info.get('ИнициалыИмениОтчестваОтветственногоЗаВУ', ''),
        '<Фамилия ответственного за ВУ>': military_info.get('ФамилияОтветсвтенногоЗаВУ', ''),
        '<Должность руководителя>': org_info.get('ДолжностьРуководителя', ''),
        '<Инициалы имени, отчества руководителя>': get_name_patronymic_initials(org_info.get('Руководитель', '')),
        '<Фамилия руководителя>': get_fio_part(org_info.get('Руководитель', ''), 0),
        '<Вид образования>': personal_info.get('Образование', {}).get('ВидОбразования', ''),
        '<Учебное заведение>': personal_info.get('Образование', {}).get('УчебноеЗаведение', ''),
        '<Квалификация>': personal_info.get('Образование', {}).get('Квалификация', ''),
        '<Специальность>': personal_info.get('Образование', {}).get('Специальность', ''),
        '<Период обучения по>': personal_info.get('Образование', {}).get('ПериодОбучения', ''),
        '<Серия квалификационного документа>': personal_info.get('Образование', {}).get('Серия', ''),
        '<Номер квалификационного документа>': personal_info.get('Образование', {}).get('Номер', ''),
        '<Серия документа, удостоверяющего личность>': document_info.get('Серия', ''),
        '<Номер документа, удостоверяющего личность>': document_info.get('Номер', ''),
        '<Дата выдачи документа, удостоверяющего личность>': document_info.get('ДатаВыдачи', ''),
        '<Орган, выдавший документ, удостоверяющего личность>': document_info.get('КемВыдан', ''),
        '<Код подразделения>': document_info.get('КодПодразделения', ''),
        '<Дата регистрации>': personal_info.get('ДатаРегистрации', ''),
        '<Инициалы имени, отчества>': get_name_patronymic_initials(personal_info.get('ФИО', '')),
        '<Водительское удостоверение:серия>': driver_license.get('Серия', ''),
        '<Водительское удостоверение:номер>': driver_license.get('Номер', ''),
        '<Водительское удостоверение:дата выдачи>': driver_license.get('ДатаВыдачи', ''),
        '<Водительское удостоверение:категории ТС>': driver_license.get('КатегорииТС', ''),
    }
    shared_input_words.update(_get_family_input_words(parsed_data_1c))
    shared_input_words.update(_commission_input_words(parsed_data_1c))
    return shared_input_words


def _replace_placeholders_in_worksheet(ws, replacements: dict) -> None:
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                for placeholder, replacement in replacements.items():
                    if placeholder in cell.value:
                        cell.value = cell.value.replace(placeholder, replacement)


def _commission_input_words(parsed_data_1c: dict) -> dict:
    org_info = parsed_data_1c['job_info']['Организация']
    return {
        '<Должность председателя комиссии>': org_info.get('ДолжностьПредседателяКомиссии', ''),
        '<Фамилия председателя комиссии>': org_info.get('ФамилияПредседателяКомиссии', ''),
        '<Инициалы имени, отчества председателя комиссии>': org_info.get('ИнициалыИмениОтчестваПредседателяКомиссии', ''),
        '<Должность члена комиссии 1>': org_info.get('ДолжностьЧленаКомиссии1', ''),
        '<Фамилия члена комиссии 1>': org_info.get('ФамилияЧленаКомиссии1', ''),
        '<Инициалы имени, отчества члена комиссии 1>': org_info.get('ИнициалыИмениОтчестваЧленаКомиссии1', ''),
        '<Должность члена комиссии 2>': org_info.get('ДолжностьЧленаКомиссии2', ''),
        '<Фамилия члена комиссии 2>': org_info.get('ФамилияЧленаКомиссии2', ''),
        '<Инициалы имени, отчества члена комиссии 2>': org_info.get('ИнициалыИмениОтчестваЧленаКомиссии2', ''),
    }


def _military_document_input_words(parsed_data_1c: dict) -> dict:
    military_info = parsed_data_1c.get('military_info', {})
    return {
        '<Адрес военкомата>': military_info.get('АдресВоенкомата', ''),
        '<Документ воинского учёта>': military_info.get('ДокументВоинскогоУчета', ''),
        '<Серия документа воинского учёта>': military_info.get('СерияДокументаВоинскогоУчета', ''),
        '<Номер документа воинского учёта>': military_info.get('НомерДокументаВоинскогоУчета', ''),
        '<Отдел, выдавший документ воинского учёта>': military_info.get('ОтделВыдавшийДокументВоинскогоУчета', ''),
        '<Дата выдачи документа воинского учёта>': military_info.get('ДатаВыдачиДокументаВоинскогоУчета', ''),
    }


parsed_data_1c =\
    {'document_info':
         {'ВидДокумента': 'Паспорт гражданина РФ',
          'ДатаВыдачи': '06.10.2021',
          'ДатаРождения': '19.09.2001',
          'КемВыдан': 'МВД по Республике Мордовия',
          'КодПодразделения': '130-001',
          'МестоРождения': 'г. СаранскРеспублика Мордовия',
          'Номер': '580748',
          'Серия': '89 21',
          'СрокДокумента': '01.01.0100'},

     'job_info':
         {'АвансПериод': 25,
          'ГрафикРаботы': 'Сменный график 1 ЛОТ',
          'ДатаПриемаНаРаботу': '10.04.2025',
          'Должность': 'Повар',
          'ЗарплатаПериод': 10,
          'ИспытательныйСрок': 3,
          'Оклад': 47581,
          'Организация':
              {'БикБанка': '',
               'ДолжностьРуководителя': 'Генеральный директор',
               'ИНН_Орг': '7704360757',
               'КПП_Орг': '770401001',
               'КорСчет': '',
               'НаименованиеПолное': 'Общество с ограниченной ответственностью "БЕЛУГА"',
               'НаименованиеСокращенное': 'ООО "БЕЛУГА"',
               'ОГРН(ИП)': '1167746557285',
               'РасСчет': '',
               'Руководитель': 'Дынов Сергей Алексеевич',
               'Телефон': '74951337214',
               'ФактическийАдрес': '125009, Москва г, Моховая ул, дом 15/1, строение 1',
               'ЮридическийАдрес': '121099, Москва г, ул Новый Арбат, д. 30/9, этаж 1, комн.2, оф.30'},

          'ОснованиеПодписи': '',
          'Отпуск': '28',
          'Подразделение': 'Горячий цех',
          'ТкНомер': '16-тр',
          'ФИО': 'Ошкин Виктор Иванович'},

     'military_info':
         {'АдресВоенкомата': '',
          'ВУС': '837037А',
          'Военкомат': 'ВК Кузьминского р-на ЮВАО г. Москвы',
          'ВоинскоеЗвание': 'Ефрейтор',
          'Годность': 'Б - годен к военной службе с незначительными ограничениями',
          'ДолжностьОтветственногоЗаВУ': 'Генеральный директор',
          'ИнициалыИмениОтчестваОтветственногоЗаВУ': 'С А',
          'КатегорияЗапаса': '1',
          'СоставПрофиль': 'Солдаты',
          'ФамилияОтветсвтенногоЗаВУ': 'Дынов'},

     'personal_info': {'Адрес_проживания': 'РОССИЯ, 109439, Город Москва, б-р Есенинский, д. 9, к. 3, кв. 59',
                       'Адрес_прописка': 'РОССИЯ, 430013, Республика Мордовия, г Саранск, ул Н.Эркая, д. 34, кв. 140',
                       'АдресБанка': '',
                       'Банк': '',
                       'БикБанка': '',
                       'ДатаРождения': '19.09.2001',
                       'ИНН': '132807896905', 'ИннБанка': '',
                       'КоррСчет': '',
                       'МестоРождения': '0,г. Саранск,,Республика Мордовия,',
                       'НомерСчета': '',
                       'Образование':
                           {'ВидОбразования': '',
                            'Квалификация': '',
                            'Номер': '',
                            'ПериодОбучения': '',
                            'Серия': '',
                            'Специальность': '',
                            'УчебноеЗаведение': ''},
                       'ПОЛ': 'Мужской',
                       'РасчетныйСчет': '',
                       'СемейноеПоложение': '',
                       'ДатаРегистрации': '',
                       'Родственники':
                           {'ДатаРождения': '',
                            'Имя': '',
                            'Отчество': '',
                            'СтепеньРодства': '',
                            'Фамилия': ''},
                       'ВодительскоеУдостоверение':
                           {'ДатаВыдачи': '',
                            'КатегорииТС': '',
                            'Номер': '',
                            'Серия': ''},
                       'СНИЛС': '169-318-152 89',
                       'Страна': 'РОССИЯ',
                       'Телефон': '79999090907',
                       'ФИО': 'Ошкин Виктор Иванович'}}


def list_msg_sample_former(input_sample: str, output_result: str, parsed_data_1c: dict) -> str:
    list_msg_input_words: dict = {
        '<Пол работника>': parsed_data_1c['personal_info'].get('ПОЛ', ''),
    }
    list_msg_input_words.update(_shared_military_registration_input_words(parsed_data_1c))

    try:
        wb = load_workbook(input_sample)
        ws = wb['Листок сообщения']
        _replace_placeholders_in_worksheet(ws, list_msg_input_words)

        wb.save(output_result)
        return output_result
    except Exception as e:
        return f'Ошибка во время формирования файла: {e}'


def msg_13_sample_former(input_sample: str, output_result: str, parsed_data_1c: dict) -> str:
    list_msg_input_words: dict = {
        '<Пол работника>': parsed_data_1c['personal_info'].get('ПОЛ', ''),
    }
    list_msg_input_words.update(_shared_military_registration_input_words(parsed_data_1c))

    try:
        wb = load_workbook(input_sample)
        ws = wb['Письмо 13']
        _replace_placeholders_in_worksheet(ws, list_msg_input_words)

        wb.save(output_result)
        return output_result
    except Exception as e:
        return f'Ошибка во время формирования файла: {e}'


def application_sample_former(input_sample: str, output_result: str, parsed_data_1c: dict, sheet_name: str) -> str:
    family_input_words = _get_family_input_words(parsed_data_1c)
    application_input_words: dict = {
        '<Краткое наименование>': parsed_data_1c['job_info']['Организация']['НаименованиеСокращенное'],
        '<ИНН организации>': parsed_data_1c['job_info']['Организация']['ИНН_Орг'],
        '<КПП организации>': parsed_data_1c['job_info']['Организация']['КПП_Орг'],
        '<ОГРН организации>': parsed_data_1c['job_info']['Организация']['ОГРН(ИП)'],
        '<Военкомат>': parsed_data_1c['military_info']['Военкомат'],

        '<Адрес военкомата>': '<Адрес военкомата>',

        '<Полное наименование>': parsed_data_1c['job_info']['Организация']['НаименованиеПолное'],
        '<Юридический адрес>': parsed_data_1c['job_info']['Организация']['ЮридическийАдрес'],

        '<Должность ответственного за ВУ>': parsed_data_1c['military_info']['ДолжностьОтветственногоЗаВУ'],
        '<Инициалы имени, отчества ответственного за ВУ>': parsed_data_1c['military_info']['ИнициалыИмениОтчестваОтветственногоЗаВУ'],
        '<Фамилия ответственного за ВУ>': parsed_data_1c['military_info']['ФамилияОтветсвтенногоЗаВУ'],

        '<Фамилия>': get_fio_part(parsed_data_1c['personal_info']['ФИО'], 0),
        '<Имя>': get_fio_part(parsed_data_1c['personal_info']['ФИО'], 1),
        '<Отчество>': get_fio_part(parsed_data_1c['personal_info']['ФИО'], 2),

        '<Воинское звание>': parsed_data_1c['military_info']['ВоинскоеЗвание'],
        '<Категория запаса>': parsed_data_1c['military_info']['КатегорияЗапаса'],
        '<Состав (профиль)>': parsed_data_1c['military_info']['СоставПрофиль'],
        '<ВУС>': parsed_data_1c['military_info']['ВУС'],
        '<Годность к воинской службе>': parsed_data_1c['military_info']['Годность'],

        '<Дата рождения>': parsed_data_1c['document_info']['ДатаРождения'],
        '<Место рождения>': parsed_data_1c['document_info']['МестоРождения'],

        '<Вид образования>': parsed_data_1c['personal_info']['Образование']['ВидОбразования'],
        '<Учебное заведение>': parsed_data_1c['personal_info']['Образование']['УчебноеЗаведение'],
        '<Квалификация>': parsed_data_1c['personal_info']['Образование']['Квалификация'],
        '<Специальность>': parsed_data_1c['personal_info']['Образование']['Специальность'],
        '<Период обучения по>': parsed_data_1c['personal_info']['Образование']['ПериодОбучения'],
        '<Семейное положение>': parsed_data_1c['personal_info'].get('СемейноеПоложение', ''),
        '<Адрес места проживания>': parsed_data_1c['personal_info']['Адрес_проживания'],
        '<Адрес по прописке>':parsed_data_1c['personal_info']['Адрес_прописка'],
        '<Подразделение>': parsed_data_1c['job_info']['Подразделение'],
        '<Должность работника>': parsed_data_1c['job_info']['Должность'],
        '<Инициалы имени, отчества руководителя>': get_name_patronymic_initials(parsed_data_1c['job_info']['Организация']['Руководитель']),
        '<Фамилия руководителя>': get_fio_part(parsed_data_1c['job_info']['Организация']['Руководитель'], 0)
    }
    application_input_words.update(_shared_military_registration_input_words(parsed_data_1c))

    try:
        wb = load_workbook(input_sample)
        ws = wb[sheet_name]
        _replace_placeholders_in_worksheet(ws, application_input_words)

        wb.save(output_result)
        return output_result
    except Exception as e:
        return f'Ошибка во время формирования файла: {e}'


def info_changes_sample_former(input_sample: str, output_result: str, parsed_data_1c: dict) -> str:
    info_changes_input_words: dict = {
        '<Краткое наименование>': parsed_data_1c['job_info']['Организация']['НаименованиеСокращенное'],
        '<ИНН организации>': parsed_data_1c['job_info']['Организация']['ИНН_Орг'],
        '<КПП организации>': parsed_data_1c['job_info']['Организация']['КПП_Орг'],
        '<ОГРН организации>': parsed_data_1c['job_info']['Организация']['ОГРН(ИП)'],
        '<Военкомат>': parsed_data_1c['military_info']['Военкомат'],

        '<Адрес военкомата>': '<Адрес военкомата>',

        '<Полное наименование>': parsed_data_1c['job_info']['Организация']['НаименованиеПолное'],
        '<Юридический адрес>': parsed_data_1c['job_info']['Организация']['ЮридическийАдрес'],

        '<Должность ответственного за ВУ>': parsed_data_1c['military_info']['ДолжностьОтветственногоЗаВУ'],
        '<Инициалы имени, отчества ответственного за ВУ>': parsed_data_1c['military_info']['ИнициалыИмениОтчестваОтветственногоЗаВУ'],
        '<Фамилия ответственного за ВУ>': parsed_data_1c['military_info']['ФамилияОтветсвтенногоЗаВУ'],

        '<Фамилия>': get_fio_part(parsed_data_1c['personal_info']['ФИО'], 0),
        '<Имя>': get_fio_part(parsed_data_1c['personal_info']['ФИО'], 1),
        '<Отчество>': get_fio_part(parsed_data_1c['personal_info']['ФИО'], 2),

        '<Воинское звание>': parsed_data_1c['military_info']['ВоинскоеЗвание'],

        '<Дата рождения>': parsed_data_1c['document_info']['ДатаРождения'],
        '<Инициалы имени, отчества руководителя>': get_name_patronymic_initials(parsed_data_1c['job_info']['Организация']['Руководитель']),
        '<Фамилия руководителя>': get_fio_part(parsed_data_1c['job_info']['Организация']['Руководитель'], 0)
    }
    info_changes_input_words.update(_shared_military_registration_input_words(parsed_data_1c))

    try:
        wb = load_workbook(input_sample)
        ws = wb['Сведения об изменениях']
        _replace_placeholders_in_worksheet(ws, info_changes_input_words)

        wb.save(output_result)
        return output_result
    except Exception as e:
        return f'Ошибка во время формирования файла: {e}'


def details_sample_former(input_sample: str, output_result: str, parsed_data_1c: dict) -> str:
    details_input_words: dict = {
        '<Краткое наименование>': parsed_data_1c['job_info']['Организация']['НаименованиеСокращенное'],
        '<ИНН организации>': parsed_data_1c['job_info']['Организация']['ИНН_Орг'],
        '<КПП организации>': parsed_data_1c['job_info']['Организация']['КПП_Орг'],
        '<ОГРН организации>': parsed_data_1c['job_info']['Организация']['ОГРН(ИП)'],
        '<Военкомат>': parsed_data_1c['military_info']['Военкомат'],

        '<Адрес военкомата>': '<Адрес военкомата>',

        '<Полное имя работника>': parsed_data_1c['personal_info']['ФИО'],
        '<Воинское звание>': parsed_data_1c['military_info']['ВоинскоеЗвание'],
        '<Дата рождения>': parsed_data_1c['document_info']['ДатаРождения'],
        '<Серия документа, удостоверяющего личность>': parsed_data_1c['document_info']['Серия'],
        '<Номер документа, удостоверяющего личность>': parsed_data_1c['document_info']['Номер'],
        '<СНИЛС>': parsed_data_1c['personal_info']['СНИЛС'],
        '<Адрес по прописке>': parsed_data_1c['personal_info']['Адрес_прописка'],
        '<Адрес места проживания>': parsed_data_1c['personal_info']['Адрес_проживания'],
        '<Полное наименование>': parsed_data_1c['job_info']['Организация']['НаименованиеПолное'],
        '<Юридический адрес>': parsed_data_1c['job_info']['Организация']['ЮридическийАдрес'],
        '<Фактический адрес>': parsed_data_1c['job_info']['Организация']['ФактическийАдрес'],
        '<Должность работника>': parsed_data_1c['job_info']['Должность'],
        '<Договор №>': parsed_data_1c['job_info']['ТкНомер'],
        '<Приказ №>': parsed_data_1c['job_info'].get('НомерПриказаОПриеме', parsed_data_1c['job_info']['ТкНомер']),
        '<Дата приказа о приеме>': parsed_data_1c['job_info']['ДатаПриемаНаРаботу'],
        '<Должность руководителя>': parsed_data_1c['job_info']['Организация']['ДолжностьРуководителя'],
        '<Инициалы имени, отчества руководителя>': get_name_patronymic_initials(parsed_data_1c['job_info']['Организация']['Руководитель']),
        '<Фамилия руководителя>': get_fio_part(parsed_data_1c['job_info']['Организация']['Руководитель'], 0),

        '<Должность ответственного за ВУ>': parsed_data_1c['military_info']['ДолжностьОтветственногоЗаВУ'],
        '<Инициалы имени, отчества ответственного за ВУ>': parsed_data_1c['military_info']['ИнициалыИмениОтчестваОтветственногоЗаВУ'],
        '<Фамилия ответственного за ВУ>': parsed_data_1c['military_info']['ФамилияОтветсвтенногоЗаВУ']
    }
    details_input_words.update(_shared_military_registration_input_words(parsed_data_1c))

    try:
        wb = load_workbook(input_sample)
        ws = wb['Сведения']
        _replace_placeholders_in_worksheet(ws, details_input_words)

        wb.save(output_result)
        return output_result
    except Exception as e:
        return f'Ошибка во время формирования файла: {e}'


def notification_sample_former(input_sample: str, output_result: str, parsed_data_1c: dict, sheet_name: str) -> str:
    notification_input_words: dict = {
        '<Полное наименование>': parsed_data_1c['job_info']['Организация']['НаименованиеПолное'],
        '<Краткое наименование>': parsed_data_1c['job_info']['Организация']['НаименованиеСокращенное'],
        '<Юридический адрес>': parsed_data_1c['job_info']['Организация']['ЮридическийАдрес'],
        '<ИНН организации>': parsed_data_1c['job_info']['Организация']['ИНН_Орг'],
        '<КПП организации>': parsed_data_1c['job_info']['Организация']['КПП_Орг'],
        '<ОГРН организации>': parsed_data_1c['job_info']['Организация']['ОГРН(ИП)'],
        '<Полное имя работника>': parsed_data_1c['personal_info']['ФИО'],

        '<Военкомат>': parsed_data_1c['military_info']['Военкомат'],

        '<Адрес военкомата>': '<Адрес военкомата>',

        '<Инициалы имени, отчества ответственного за ВУ>': parsed_data_1c['military_info']['ИнициалыИмениОтчестваОтветственногоЗаВУ'],
        '<Фамилия ответственного за ВУ>': parsed_data_1c['military_info']['ФамилияОтветсвтенногоЗаВУ']
    }
    notification_input_words.update(_shared_military_registration_input_words(parsed_data_1c))

    try:
        wb = load_workbook(input_sample)
        ws = wb[sheet_name]
        _replace_placeholders_in_worksheet(ws, notification_input_words)

        wb.save(output_result)
        return output_result
    except Exception as e:
        return f'Ошибка во время формирования файла: {e}'


def f10_sample_former(input_sample: str, output_result: str, parsed_data_1c: dict) -> str:
    family_input_words = _get_family_input_words(parsed_data_1c)
    f10_input_words: dict = {
        '<Полное наименование>': parsed_data_1c['job_info']['Организация']['НаименованиеПолное'],

        '<Инициалы имени, отчества ответственного за ВУ>': parsed_data_1c['military_info']['ИнициалыИмениОтчестваОтветственногоЗаВУ'],
        '<Фамилия ответственного за ВУ>': parsed_data_1c['military_info']['ФамилияОтветсвтенногоЗаВУ'],

        '<Фамилия>': get_fio_part(parsed_data_1c['personal_info']['ФИО'], 0),
        '<Имя>': get_fio_part(parsed_data_1c['personal_info']['ФИО'], 1),
        '<Отчество>': get_fio_part(parsed_data_1c['personal_info']['ФИО'], 2),
        '<Дата рождения>': parsed_data_1c['document_info']['ДатаРождения'],
        '<Место рождения>': parsed_data_1c['document_info']['МестоРождения'],

        '<Вид образования>': parsed_data_1c['personal_info']['Образование']['ВидОбразования'],
        '<Учебное заведение>': parsed_data_1c['personal_info']['Образование']['УчебноеЗаведение'],
        '<Квалификация>': parsed_data_1c['personal_info']['Образование']['Квалификация'],
        '<Специальность>': parsed_data_1c['personal_info']['Образование']['Специальность'],
        '<Серия квалификационного документа>': parsed_data_1c['personal_info']['Образование']['Серия'],
        '<Номер квалификационного документа>': parsed_data_1c['personal_info']['Образование']['Номер'],
        '<Период обучения по>': parsed_data_1c['personal_info']['Образование']['ПериодОбучения'],

        '<Серия документа, удостоверяющего личность>': parsed_data_1c['document_info']['Серия'],
        '<Номер документа, удостоверяющего личность>': parsed_data_1c['document_info']['Номер'],
        '<Дата выдачи документа, удостоверяющего личность>': parsed_data_1c['document_info']['ДатаВыдачи'],
        '<Орган, выдавший документ, удостоверяющего личность>': parsed_data_1c['document_info']['КемВыдан'],
        '<Код подразделения>': parsed_data_1c['document_info']['КодПодразделения'],
        '<Адрес по прописке>': parsed_data_1c['personal_info']['Адрес_прописка'],

        '<Дата регистрации>': parsed_data_1c['personal_info'].get('ДатаРегистрации', ''),
        '<Семейное положение>': parsed_data_1c['personal_info'].get('СемейноеПоложение', ''),

        '<Адрес места проживания>': parsed_data_1c['personal_info']['Адрес_проживания'],
        '<Мобильный телефон>': parsed_data_1c['personal_info']['Телефон'],

        '<Категория запаса>': parsed_data_1c['military_info']['КатегорияЗапаса'],
        '<Военкомат>': parsed_data_1c['military_info']['Военкомат'],
        '<Воинское звание>': parsed_data_1c['military_info']['ВоинскоеЗвание'],
        '<Состав (профиль)>': parsed_data_1c['military_info']['СоставПрофиль'],
        '<ВУС>': parsed_data_1c['military_info']['ВУС'],
        '<Годность к воинской службе>': parsed_data_1c['military_info']['Годность'],

        '<Документ воинского учёта>': '<Документ воинского учёта>',
        '<Серия документа воинского учёта>': '<Серия документа воинского учёта>',
        '<Номер документа воинского учёта>': '<Номер документа воинского учёта>',
        '<Отдел, выдавший документ воинского учёта>': '<Отдел, выдавший документ воинского учёта>',
        '<Дата выдачи документа воинского учёта>': '<Дата выдачи документа воинского учёта>',

        '<Номер приказа о приеме>': parsed_data_1c['job_info'].get('НомерПриказаОПриеме', parsed_data_1c['job_info']['ТкНомер']),
        '<Приказ №>': parsed_data_1c['job_info'].get('НомерПриказаОПриеме', parsed_data_1c['job_info']['ТкНомер']),
        '<Дата приказа о приеме>': parsed_data_1c['job_info']['ДатаПриемаНаРаботу'],
        '<Должность работника>': parsed_data_1c['job_info']['Должность'],

        '<Должность ответственного за ВУ>': parsed_data_1c['military_info']['ДолжностьОтветственногоЗаВУ'],

        '<Инициалы имени, отчества>': get_name_patronymic_initials(parsed_data_1c['personal_info']['ФИО']),
        '<Водительское удостоверение:серия>': parsed_data_1c['personal_info'].get('ВодительскоеУдостоверение', {}).get('Серия', ''),
        '<Водительское удостоверение:номер>': parsed_data_1c['personal_info'].get('ВодительскоеУдостоверение', {}).get('Номер', ''),
        '<Водительское удостоверение:дата выдачи>': parsed_data_1c['personal_info'].get('ВодительскоеУдостоверение', {}).get('ДатаВыдачи', ''),
        '<Водительское удостоверение:категории ТС>': parsed_data_1c['personal_info'].get('ВодительскоеУдостоверение', {}).get('КатегорииТС', '')
    }
    f10_input_words.update(_shared_military_registration_input_words(parsed_data_1c))

    try:
        wb = load_workbook(input_sample)
        ws = wb['Карточка']
        _replace_placeholders_in_worksheet(ws, f10_input_words)


        wb.save(output_result)
        return output_result
    except Exception as e:
        return f'Ошибка во время формирования файла: {e}'


# def form_military_files(input_folder: str, output_result: str, parsed_data_1c: dict) -> list[str]:
#     formed_files: list[str] = []
#
#     for file in os.listdir(input_folder):
#         list_msg: str = list_msg_sample_former(input_sample=f'{input_folder}/{file}',
#                                                output_result=f'{output_result}/output_{file}',
#                                                parsed_data_1c=parsed_data_1c)
#
#         # Значит файл сформирован без возникновения ошибок
#         if 'output' in list_msg:
#             formed_files.append(list_msg)
#
#         msg_13: str = msg_13_sample_former(input_sample=f'{input_folder}/{file}',
#                                            output_result=f'{output_result}/output_{file}',
#                                            parsed_data_1c=parsed_data_1c)
#
#     return formed_files




if __name__ == '__main__':
    list_msg: str = list_msg_sample_former(input_sample='input_samples/Листок-сообщения.xlsx',
                                           output_result='output_result/output_Листок-сообщения.xlsx',
                                           parsed_data_1c=parsed_data_1c)
    #
    # msg_13: str = msg_13_sample_former(input_sample='input_samples/Письмо 13.xlsx',
    #                                    output_result='output_result/output_Письмо 13.xlsx',
    #                                    parsed_data_1c=parsed_data_1c)

    # appl_10: str = application_sample_former(input_sample='input_samples/Приложение 10.xlsx',
    #                                          output_result='output_result/output_Приложение 10.xlsx',
    #                                          parsed_data_1c=parsed_data_1c,
    #                                          sheet_name='Сведения прил. 10')

    # appl_23: str = application_sample_former(input_sample='input_samples/Приложение 23.xlsx',
    #                                          output_result='output_result/output_Приложение 23.xlsx',
    #                                          parsed_data_1c=parsed_data_1c,
    #                                          sheet_name='Сверка')
    #
    # info_change: str = info_changes_sample_former(input_sample='input_samples/Сведения об изменениях.xlsx',
    #                                               output_result='output_result/output_Сведения об изменениях.xlsx',
    #                                               parsed_data_1c=parsed_data_1c)
    #
    # details: str = details_sample_former(input_sample='input_samples/Сведения.xlsx',
    #                                      output_result='output_result/output_Сведения.xlsx',
    #                                      parsed_data_1c=parsed_data_1c)
    #
    # note_1: str = notification_sample_former(input_sample='input_samples/Уведомление 1.xlsx',
    #                                          output_result='output_result/output_Уведомление 1.xlsx',
    #                                          parsed_data_1c=parsed_data_1c, sheet_name='Уведомление 1')
    #
    # note_2: str = notification_sample_former(input_sample='input_samples/Уведомление 2.xlsx',
    #                                          output_result='output_result/output_Уведомление 2.xlsx',
    #                                          parsed_data_1c=parsed_data_1c, sheet_name='Уведомление 2')
    #

    # f10_doc: str = f10_sample_former(input_sample='input_samples/Ф10.xlsx',
    #                                  output_result='output_result/output_Ф10.xlsx',
    #                                  parsed_data_1c=parsed_data_1c)
