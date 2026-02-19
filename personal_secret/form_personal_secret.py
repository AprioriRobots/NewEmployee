from docx import Document
from python_docx_replace import docx_replace, docx_get_keys


# def personal_secret_sample_former(input_sample: str, output_result: str, parsed_data_1c: dict) -> str:
#     PERSONAL_SECRET_INPUT_WORDS: dict = {
#         'Дата приказа о приеме': parsed_data_1c['job_info']['ДатаПриемаНаРаботу'],
#         'Полное имя работника': parsed_data_1c['personal_info']['ФИО'],
#         'Вид документа, удостоверяющего личность': parsed_data_1c['document_info']['ВидДокумента'],
#         'Серия документа, удостоверяющего личность': parsed_data_1c['document_info']['Серия'],
#         'Номер документа, удостоверяющего личность': parsed_data_1c['document_info']['Серия'],
#         'Дата выдачи документа, удостоверяющего личность': parsed_data_1c['document_info']['ДатаВыдачи'],
#         'Орган, выдавший документ, удостоверяющего личность': parsed_data_1c['document_info']['КемВыдан'],
#         'Номер телефона работника': parsed_data_1c['personal_info']['Телефон'],
#         'Адрес по прописке': parsed_data_1c['personal_info']['Адрес_прописка'],
#         'Полное наименование':  parsed_data_1c['job_info']['Организация']['НаименованиеПолное'],
#         'Краткое наименование': parsed_data_1c['job_info']['Организация']['НаименованиеСокращенное'],
#         'ИНН организации': parsed_data_1c['job_info']['Организация']['ИНН_Орг'],
#         'ОГРН организации': parsed_data_1c['job_info']['Организация']['ОГРН(ИП)'],
#         'Должность руководителя': parsed_data_1c['job_info']['Организация']['ДолжностьРуководителя'],
#         'Руководитель': parsed_data_1c['job_info']['Организация']['Руководитель'],
#
#         'Основания подписи': 'Основания подписи',
#
#         'Фактический адрес': parsed_data_1c['job_info']['Организация']['ФактическийАдрес'],
#         'Банк работника': parsed_data_1c['personal_info']['БикБанка'],
#         'Адрес банка работника': parsed_data_1c['personal_info']['АдресБанка'],
#         'ИНН банка работника': parsed_data_1c['personal_info']['ИннБанка']
#     }
#
#     try:
#         doc = Document(input_sample)
#
#         paragraph_count: int = len(doc.paragraphs)
#         table_count: int = len(doc.tables)
#
#         # Итерируемся по всем таблицам в документе
#         for i in range(table_count):
#             # Итерируемся по каждой строке
#             for row in doc.tables[i].table.rows:
#                 current_row = row
#                 cell_count_of_current_row: int = len(current_row.cells)
#
#                 # Итерируемся по каждой ячейке
#                 for j in range(cell_count_of_current_row):
#                     current_cell_txt = current_row.cells[j].text
#
#                     if '<Банк работника>' in current_cell_txt:
#                         update_text = current_cell_txt.replace('<Банк работника>', PERSONAL_SECRET_INPUT_WORDS['Банк_работника'])
#
#                         current_row.cells[j].text = update_text
#                         current_cell_txt = update_text
#
#                     if '<Адрес банка работника>' in current_cell_txt:
#                         update_text = current_cell_txt.replace('<Адрес банка работника>',
#                                                                PERSONAL_SECRET_INPUT_WORDS['Адрес_банка_работника'])
#                         current_row.cells[j].text = update_text
#                         current_cell_txt = update_text
#
#                     if '<ИНН банка работника>' in current_cell_txt:
#                         update_text = current_cell_txt.replace('<ИНН банка работника>',
#                                                                PERSONAL_SECRET_INPUT_WORDS['ИНН_банка_работника'])
#                         current_row.cells[j].text = update_text
#                         current_cell_txt = update_text
#
#         for i in range(paragraph_count):
#             current_txt = doc.paragraphs[i].text
#
#             if '<Дата приказа о приеме>' in current_txt:
#                 updated_text = current_txt.replace('<Дата приказа о приеме>',
#                                                    PERSONAL_SECRET_INPUT_WORDS['Дата_приказа_о_приеме'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Полное имя работника>' in current_txt:
#                 updated_text = current_txt.replace('<Полное имя работника>',
#                                                    PERSONAL_SECRET_INPUT_WORDS['Полное_имя_работника'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Вид документа, удостоверяющего личность>' in current_txt:
#                 updated_text = current_txt.replace('<Вид документа, удостоверяющего личность>',
#                                                    PERSONAL_SECRET_INPUT_WORDS['Вид_документа_удостоверяющего_личность'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Серия документа, удостоверяющего личность>' in current_txt:
#                 updated_text = current_txt.replace('<Серия документа, удостоверяющего личность>',
#                                                    PERSONAL_SECRET_INPUT_WORDS['Серия_документа_удостоверяющего_личность'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Номер документа, удостоверяющего личность>' in current_txt:
#                 updated_text = current_txt.replace('<Номер документа, удостоверяющего личность>',
#                                                    PERSONAL_SECRET_INPUT_WORDS['Номер_документа_удостоверяющего_личность'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Дата выдачи документа, удостоверяющего личность>' in current_txt:
#                 updated_text = current_txt.replace(
#                     '<Дата выдачи документа, удостоверяющего личность>',
#                     PERSONAL_SECRET_INPUT_WORDS['Дата_выдачи_документа_удостоверяющего_личность'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Орган, выдавший документ, удостоверяющего личность>' in current_txt:
#                 updated_text = current_txt.replace(
#                     '<Орган, выдавший документ, удостоверяющего личность>',
#                     PERSONAL_SECRET_INPUT_WORDS['Орган_выдавший_документ_удостоверяющего_личность'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Номер телефона работника>' in current_txt:
#                 updated_text = current_txt.replace('<Номер телефона работника>',
#                                                    PERSONAL_SECRET_INPUT_WORDS['Номер_телефона_работника'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Адрес по прописке>' in current_txt:
#                 updated_text = current_txt.replace('<Адрес по прописке>',
#                                                    PERSONAL_SECRET_INPUT_WORDS['Адрес_по_прописке'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Полное наименование>' in current_txt:
#                 updated_text = current_txt.replace('<Полное наименование>',
#                                                    PERSONAL_SECRET_INPUT_WORDS['Полное_наименование'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Краткое наименование>' in current_txt:
#                 updated_text = current_txt.replace('<Краткое наименование>',
#                                                    PERSONAL_SECRET_INPUT_WORDS['Краткое_наименование'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<ИНН организации>' in current_txt:
#                 updated_text = current_txt.replace('<ИНН организации>',
#                                                    PERSONAL_SECRET_INPUT_WORDS['ИНН_организации'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<ОГРН организации>' in current_txt:
#                 updated_text = current_txt.replace('<ОГРН организации>',
#                                                    PERSONAL_SECRET_INPUT_WORDS['ОГРН_организации'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Должность руководителя>' in current_txt:
#                 updated_text = current_txt.replace('<Должность руководителя>',
#                                                    PERSONAL_SECRET_INPUT_WORDS['Должность_руководителя'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Руководитель>' in current_txt:
#                 updated_text = current_txt.replace('<Руководитель>',
#                                                    PERSONAL_SECRET_INPUT_WORDS['Руководитель'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Основания подписи>' in current_txt:
#                 updated_text = current_txt.replace('<Основания подписи>',
#                                                    PERSONAL_SECRET_INPUT_WORDS['Основания_подписи'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Фактический адрес>' in current_txt:
#                 updated_text = current_txt.replace('<Фактический адрес>',
#                                                    PERSONAL_SECRET_INPUT_WORDS['Фактический_адрес'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#         doc.save(output_result)
#         return f'Файл успешно сформирован. Сохранен в {output_result}'
#
#     except Exception as e:
#         return f'Ошибка во время формирования файла: {e}'


def personal_secret_sample_former(input_sample: str, output_result: str, parsed_data_1c: dict) -> str:
    personal_secret_input_words: dict = {
        'Дата приказа о приеме': parsed_data_1c['job_info']['ДатаПриемаНаРаботу'],
        'Полное имя работника': parsed_data_1c['personal_info']['ФИО'],
        'Вид документа, удостоверяющего личность': parsed_data_1c['document_info']['ВидДокумента'],
        'Серия документа, удостоверяющего личность': parsed_data_1c['document_info']['Серия'],
        'Номер документа, удостоверяющего личность': parsed_data_1c['document_info']['Серия'],
        'Дата выдачи документа, удостоверяющего личность': parsed_data_1c['document_info']['ДатаВыдачи'],
        'Орган, выдавший документ, удостоверяющего личность': parsed_data_1c['document_info']['КемВыдан'],
        'Номер телефона работника': parsed_data_1c['personal_info']['Телефон'],
        'Адрес по прописке': parsed_data_1c['personal_info']['Адрес_прописка'],
        'Полное наименование':  parsed_data_1c['job_info']['Организация']['НаименованиеПолное'],
        'Краткое наименование': parsed_data_1c['job_info']['Организация']['НаименованиеСокращенное'],
        'ИНН организации': parsed_data_1c['job_info']['Организация']['ИНН_Орг'],
        'ОГРН организации': parsed_data_1c['job_info']['Организация']['ОГРН(ИП)'],
        'Должность руководителя': parsed_data_1c['job_info']['Организация']['ДолжностьРуководителя'],
        'Руководитель': parsed_data_1c['job_info']['Организация']['Руководитель'],

        'Основания подписи': 'Основания подписи',

        'Фактический адрес': parsed_data_1c['job_info']['Организация']['ФактическийАдрес'],
        'Банк работника': parsed_data_1c['personal_info']['БикБанка'],
        'Адрес банка работника': parsed_data_1c['personal_info']['АдресБанка'],
        'ИНН банка работника': parsed_data_1c['personal_info']['ИннБанка']
    }

    try:
        doc = Document(input_sample)
        keys = docx_get_keys(doc, pattern='<>')
        print(keys)
        docx_replace(doc, '<>', **personal_secret_input_words)

        doc.save(output_result)
        return output_result

    except Exception as e:
        return f'Ошибка во время формирования файла: {e}'


if __name__ == '__main__':
    result: str = personal_secret_sample_former(
        input_sample='input_sample/Согласие на распространение персональных данных НОВОЕ.docx',
        output_result='output_result/output_Согласие на распространение персональных данных НОВОЕ.docx'
    )

    result_2: str = personal_secret_sample_former(
        input_sample='input_sample/Согласие на обработку персональных данных НОВОЕ.docx',
        output_result='output_result/output_Согласие на обработку персональных данных НОВОЕ.docx'
    )
