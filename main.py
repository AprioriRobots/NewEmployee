from docx import Document


INPUT_WORDS: dict = {
    'Договор': '5553535',
    'Дата_приказа_о_приеме': '13.01.2026',
    'Полное_наименование': 'ООО Априори',
    'Должность_руководителя': 'Директор компании',
    'Руководитель': 'Иванов Иван Иванович',
    'Основания_подписи': 'Заключения ТК РФ',
    'Гражданство_страны': 'РФ',
    'Полное_имя_работника': 'Петров Петр Петрович',
    'Подразделение': 'Финансов',
    'Должность_работника': 'Бухгалтер',
    'Краткое_наименование': 'компания Априори',
    'Дата_приема': '14.01.2026',
    'Испытательный_срок': '3',
    'График_работы': '8 часов',
    'Длительность_рабочей_недели': '40',
    'Количество_дней_ежегодного_отпуска': '28',
    'Оклад': '35 000',
    'Дата_выплаты_аванса': '12',
    'Дата_выплаты_зарплаты': '24',
}


def tkrf_sample_updater():
    doc = Document('input_sample/Трудовой договор.docx')
    paragraph_count: int = len(doc.paragraphs)

    for table in doc.tables:
        print(table)

    for i in range(paragraph_count):
        print(doc.paragraphs[i].text)
        current_txt = doc.paragraphs[i].text

        if '<Договор №' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Договор №>', INPUT_WORDS['Договор'])
        elif '<Полное наименование>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Полное наименование>', INPUT_WORDS['Полное_наименование'])
        elif '<Должность руководителя>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Должность руководителя>',
                                                         INPUT_WORDS['Должность_руководителя'])
        elif '<Руководитель>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Руководитель>', INPUT_WORDS['Руководитель'])
        elif '<Основания подписи>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Основания подписи>', INPUT_WORDS['Основания_подписи'])
        elif '<Гражданство страны>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Гражданство страны>', INPUT_WORDS['Гражданство_страны'])
        elif '<Полное имя работника>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Полное имя работника>', INPUT_WORDS['Полное_имя_работника'])
        elif '<Подразделение>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Подразделение>', INPUT_WORDS['Подразделение'])
        elif '<Должность работника>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Должность работника>', INPUT_WORDS['Должность_работника'])
        elif '<Краткое наименование>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Краткое наименование>', INPUT_WORDS['Краткое_наименование'])
        elif '<Дата приема>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Дата приема>', INPUT_WORDS['Дата_приема'])
        elif '<Испытательный срок>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Испытательный срок>', INPUT_WORDS['Испытательный_срок'])
        elif '<График работы>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<График работы>', INPUT_WORDS['График_работы'])
        elif '<Длительность рабочей недели>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Длительность рабочей недели>',
                                                         INPUT_WORDS['Длительность_рабочей_недели'])
        elif '<Количество дней ежегодного отпуска>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Количество дней ежегодного отпуска>',
                                                         INPUT_WORDS['Количество_дней_ежегодного_отпуска'])
        elif '<Оклад>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Оклад>', INPUT_WORDS['Оклад'])
        elif '<Дата выплаты аванса>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Дата выплаты аванса>', INPUT_WORDS['Дата_выплаты_аванса'])
        elif '<Дата выплаты зарплаты>' in current_txt:
            doc.paragraphs[i].text = current_txt.replace('<Дата выплаты зарплаты>',
                                                         INPUT_WORDS['Дата_выплаты_зарплаты'])

    doc.save('output_result/output_Трудовой договор.docx')


if __name__ == '__main__':
    pass
