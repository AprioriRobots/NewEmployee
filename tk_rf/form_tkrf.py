import re
import zipfile
import xml.etree.ElementTree as ET
from copy import deepcopy
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from template_utils import replace_text_with_relative_dates

try:
    import win32com.client
except ImportError:  # pragma: no cover
    win32com = None


_PLACEHOLDER_PATTERN = re.compile(r"<[^<>]+>")
_WORDPROCESSING_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W_TAG = f"{{{_WORDPROCESSING_NS}}}"
_DEBUG_LOG_PATH = Path("tk_rf_debug.log")


def _write_tkrf_debug(message: str) -> None:
    # Contract diagnostics are disabled until they are explicitly needed again.
    return None


def _write_minor_contract_debug(message: str) -> None:
    return None


def _normalize_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip().lower().replace("ё", "е")


def _is_truthy(value) -> bool:
    return _normalize_text(value) in ("y", "yes", "true", "1", "да", "истина")


def _normalize_placeholder_name(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.startswith("<") and text.endswith(">"):
        text = text[1:-1].strip()
    return " ".join(text.split()).lower()


def _parse_date(value) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())

    text = str(value).replace("\xa0", " ").strip()
    for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _to_int(value) -> int | None:
    if value in (None, ""):
        return None
    try:
        if isinstance(value, str):
            cleaned = value.replace(" ", "").replace(",", ".")
            return int(float(cleaned))
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _choose_plural(value: int, forms: tuple[str, str, str]) -> str:
    value = abs(value) % 100
    if 11 <= value <= 19:
        return forms[2]
    last_digit = value % 10
    if last_digit == 1:
        return forms[0]
    if 2 <= last_digit <= 4:
        return forms[1]
    return forms[2]


def _number_to_words_ru(value) -> str:
    number = _to_int(value)
    if number is None:
        return ""
    if number == 0:
        return "ноль"

    ones_male = ("", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять")
    ones_female = ("", "одна", "две", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять")
    teens = (
        "десять", "одиннадцать", "двенадцать", "тринадцать", "четырнадцать",
        "пятнадцать", "шестнадцать", "семнадцать", "восемнадцать", "девятнадцать",
    )
    tens = ("", "", "двадцать", "тридцать", "сорок", "пятьдесят", "шестьдесят", "семьдесят", "восемьдесят", "девяносто")
    hundreds = ("", "сто", "двести", "триста", "четыреста", "пятьсот", "шестьсот", "семьсот", "восемьсот", "девятьсот")
    orders = (
        (("", "", ""), False),
        (("тысяча", "тысячи", "тысяч"), True),
        (("миллион", "миллиона", "миллионов"), False),
        (("миллиард", "миллиарда", "миллиардов"), False),
    )

    def triplet_to_words(triplet: int, female: bool) -> list[str]:
        parts: list[str] = []
        parts.append(hundreds[triplet // 100])
        remainder = triplet % 100
        if 10 <= remainder <= 19:
            parts.append(teens[remainder - 10])
        else:
            parts.append(tens[remainder // 10])
            ones = ones_female if female else ones_male
            parts.append(ones[remainder % 10])
        return [part for part in parts if part]

    parts: list[str] = []
    order_index = 0
    while number > 0 and order_index < len(orders):
        triplet = number % 1000
        if triplet:
            order_forms, is_female = orders[order_index]
            triplet_words = triplet_to_words(triplet, is_female)
            if order_index > 0:
                triplet_words.append(_choose_plural(triplet, order_forms))
            parts.insert(0, " ".join(triplet_words))
        number //= 1000
        order_index += 1

    return " ".join(parts).strip()


def _calculate_age_on_date(work_date: str, birth_date: str) -> int | None:
    work_dt = _parse_date(work_date)
    birth_dt = _parse_date(birth_date)
    if not work_dt or not birth_dt:
        return None

    age = work_dt.year - birth_dt.year
    if (work_dt.month, work_dt.day) < (birth_dt.month, birth_dt.day):
        age -= 1
    return age


def _build_contract_placeholders(parsed_data_1c: dict) -> dict:
    job_info = parsed_data_1c["job_info"]
    org_info = job_info["Организация"]
    personal_info = parsed_data_1c["personal_info"]
    document_info = parsed_data_1c["document_info"]
    bitrix_fields = parsed_data_1c.get("bitrix_fields", {})
    is_foreign_worker = _is_truthy(bitrix_fields.get("UF_CRM_3_1768399223", ""))
    migrant_document_type = str(job_info.get("ВидДокументаМигранта", "") or "")
    document_type = migrant_document_type or str(document_info.get("ВидДокумента", "") or "")
    document_type_normalized = _normalize_text(document_type)
    patent_checked = _is_truthy(job_info.get("ЧекбоксПатент", "")) or "патент" in document_type_normalized
    permit_checked = _is_truthy(job_info.get("ЧекбоксРазрешение", "")) or any(
        marker in document_type_normalized
        for marker in (
            "разрешение",
            "студент",
            "вид на жительство",
            "внж",
            "разрешение на временное проживание",
            "рвп",
        )
    )

    replacements = {
        "<Договор №>": job_info.get("ТкНомер", ""),
        "<Дата приказа о приеме>": job_info.get("ДатаПриемаНаРаботу", ""),
        "<Полное наименование>": org_info.get("НаименованиеПолное", ""),
        "<Должность руководителя>": org_info.get("ДолжностьРуководителя", ""),
        "<Руководитель>": org_info.get("Руководитель", ""),
        "<Основания подписи>": job_info.get("ОснованиеПодписи", ""),
        "<Гражданство страны>": personal_info.get("Страна", ""),
        "<Иностранный гражданин>": "Да" if is_foreign_worker else "Нет",
        "<Иностранный работник>": "Да" if is_foreign_worker else "Нет",
        "<Чекбокс Патент>": "Да" if patent_checked else "Нет",
        "<Чекбокс Разрешение>": "Да" if permit_checked else "Нет",
        "<Полное имя работника>": personal_info.get("ФИО", ""),
        "<Подразделение>": job_info.get("Подразделение", ""),
        "<Должность работника>": job_info.get("Должность", ""),
        "<Краткое наименование>": org_info.get("НаименованиеСокращенное", ""),
        "<Дата приема>": job_info.get("ДатаПриемаНаРаботу", ""),
        "<Испытательный срок>": str(job_info.get("ИспытательныйСрок", "")),
        "<Испытательный срок словами>": _number_to_words_ru(job_info.get("ИспытательныйСрок", "")),
        "<График работы>": job_info.get("ГрафикРаботы", ""),
        "<Длительность рабочей недели>": str(job_info.get("ДлительностьРабочейНедели", "")),
        "<Количество дней ежегодного отпуска>": str(job_info.get("Отпуск", "28")),
        "<Оклад>": str(job_info.get("Оклад", "")),
        "<Оклад словами>": _number_to_words_ru(job_info.get("Оклад", "")),
        "<Дата выплаты аванса>": str(job_info.get("АвансПериод", "")),
        "<Дата выплаты зарплаты>": str(job_info.get("ЗарплатаПериод", "")),
        "<Юридический адрес>": org_info.get("ЮридическийАдрес", ""),
        "<Фактический адрес>": org_info.get("ФактическийАдрес", ""),
        "<ИНН организации>": org_info.get("ИНН_Орг", ""),
        "<КПП организации>": org_info.get("КПП_Орг", ""),
        "<ОГРН организации>": org_info.get("ОГРН(ИП)", ""),
        "<Расчетный счет организации>": org_info.get("РасСчет", ""),
        "<БИК банка организации>": org_info.get("БикБанка", ""),
        "<Корреспондентский счет организации>": org_info.get("КорСчет", ""),
        "<Дата рождения>": document_info.get("ДатаРождения", "") or personal_info.get("ДатаРождения", ""),
        "<Место рождения>": document_info.get("МестоРождения", ""),
        "<Вид документа>": migrant_document_type,
        "<Серия>": (
            job_info.get("СерияПатента", "")
            if patent_checked
            else job_info.get("СерияРазрешения", "")
        ),
        "<Номер>": (
            job_info.get("НомерПатента", "")
            if patent_checked
            else job_info.get("НомерРазрешения", "")
        ),
        "<Начало>": (
            job_info.get("ДатаВыдачиПатента", "")
            if patent_checked
            else job_info.get("ДатаВыдачиРазрешения", "")
        ),
        "<Дата выдачи>": (
            job_info.get("ДатаВыдачиПатента", "")
            if patent_checked
            else job_info.get("ДатаВыдачиРазрешения", "")
        ),
        "<Кем выдан>": (
            job_info.get("КемВыданПатент", "")
            if patent_checked
            else job_info.get("КемВыданРазрешения", "")
        ),
        "<Серия документа, удостоверяющего личность>": document_info.get("Серия", ""),
        "<Номер документа, удостоверяющего личность>": document_info.get("Номер", ""),
        "<Орган, выдавший документ, удостоверяющего личность>": document_info.get("КемВыдан", ""),
        "<Дата выдачи документа, удостоверяющего личность>": document_info.get("ДатаВыдачи", ""),
        "<СНИЛС>": personal_info.get("СНИЛС", ""),
        "<Адрес по прописке>": personal_info.get("Адрес_прописка", ""),
        "<Рабочая дата>": job_info.get("ДатаПриемаНаРаботу", ""),
        "<Патент>": job_info.get("Патент", "") if patent_checked else "",
        "<Серия патента>": job_info.get("СерияПатента", ""),
        "<Номер патента>": job_info.get("НомерПатента", ""),
        "<Дата выдачи патента>": job_info.get("ДатаВыдачиПатента", ""),
        "<Патент выдан>": job_info.get("КемВыданПатент", ""),
        "<ИнформацияОПатенте>": "",
        "<Явка>": job_info.get("Явка", ""),
        "<Срочный договор>": job_info.get("СрочныйДоговор", ""),
        "<Срочный трудовой договор до>": job_info.get("СрочныйТрудовойДоговорДо", ""),
        "<Основание срочного договора>": job_info.get("ОснованиеСрочногоДоговора", ""),
        "<Вид занятости>": job_info.get("ВидЗанятости", ""),
        "<Оплата по часовому тарифу>": str(job_info.get("ОплатаПоЧасовомуТарифу", "")),
        "<Начисление>": job_info.get("Начисление", ""),
        "<В организации применяется КТУ>": str(job_info.get("ВОрганизацииПрименяетсяКТУ", "")),
    }
    return replacements


def _apply_rule_driven_placeholders(parsed_data_1c: dict, replacements: dict) -> dict:
    updated_replacements = dict(replacements)
    rules = _load_formulation_rules()
    patent_targets = {
        "ИнформацияОПатенте",
        'Реквизиты Работника в разделе "Адреса и реквизиты Сторон" трудового договора',
    }

    patent_texts: list[str] = []
    for rule in rules:
        point = str(rule.get("point") or "").strip()
        if point not in patent_targets:
            continue
        if not _rule_matches_with_foreign_fallback(rule, updated_replacements):
            continue

        rendered_text = _replace_placeholders(rule["text"], updated_replacements).strip()
        if rendered_text and rendered_text not in patent_texts:
            patent_texts.append(rendered_text)

    updated_replacements["<ИнформацияОПатенте>"] = "\n".join(patent_texts)
    return updated_replacements


def _is_hourly_payment(parsed_data_1c: dict, replacements: dict) -> bool:
    job_info = parsed_data_1c.get("job_info", {})
    hourly_rate = str(job_info.get("ОплатаПоЧасовомуТарифу", "") or replacements.get("<Оплата по часовому тарифу>", "") or "").strip()
    accrual_name = _normalize_text(job_info.get("Начисление", "") or replacements.get("<Начисление>", ""))
    tariff_kind = _normalize_text(job_info.get("ВидТарифнойСтавки", ""))

    if not hourly_rate:
        return False

    if any(marker in tariff_kind for marker in ("час", "почас")):
        return True

    if any(marker in accrual_name for marker in ("час", "почас", "тариф")):
        return True

    return False


def _get_hourly_payment_rule_text(rules: list[dict], replacements: dict) -> str:
    for rule in rules:
        point = str(rule.get("point") or "").strip()
        condition = str(rule.get("condition") or "")
        if point != "6.1.":
            continue
        if 'Раздел "Начисление" содержит "Оплата по часовому тарифу"' not in condition:
            continue
        return _replace_placeholders(str(rule.get("text") or "").strip(), replacements)

    return _replace_placeholders(
        (
            "Работнику за выполнение трудовых обязанностей по настоящему Договору "
            "устанавливается заработная плата в соответствии с тарифной ставкой "
            "в размере <Оплата по часовому тарифу> рублей за единицу времени (один час). "
            "Исходя из условий и режима работы начисление заработной платы производится "
            "пропорционально отработанному времени. Работодатель удерживает из суммы "
            "заработной платы Работника все налоги и сборы в соответствии с действующим "
            "законодательством РФ. Указанная заработная плата может быть повышена, "
            "проиндексирована по решению <Должность руководителя>."
        ),
        replacements,
    )


def _replace_placeholders(text: str, replacements: dict) -> str:
    result = replace_text_with_relative_dates(text, replacements)
    normalized_replacements = {
        _normalize_placeholder_name(placeholder): str(value)
        for placeholder, value in replacements.items()
    }

    def replace_match(match: re.Match) -> str:
        placeholder = match.group(0)
        if placeholder in replacements:
            return str(replacements[placeholder])
        return normalized_replacements.get(_normalize_placeholder_name(placeholder), placeholder)

    return _PLACEHOLDER_PATTERN.sub(replace_match, result)


def _load_formulation_rules() -> list[dict]:
    workbook_path = next(p for p in Path("formul_tkrf/input_sample").glob("*.xlsx") if not p.name.startswith("~$"))
    workbook = load_workbook(workbook_path, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]

    rules: list[dict] = []
    row_index = 2
    while True:
        rule_index = worksheet[f"B{row_index}"].value
        point = worksheet[f"C{row_index}"].value
        text = worksheet[f"D{row_index}"].value
        condition = worksheet[f"E{row_index}"].value
        action = worksheet[f"F{row_index}"].value

        if not any((rule_index, point, text, condition, action)):
            break

        rules.append(
            {
                "index": rule_index,
                "point": str(point or "").strip(),
                "text": str(text or "").strip(),
                "condition": str(condition or "").strip(),
                "action": str(action or "").strip(),
            }
        )
        row_index += 1

    return rules


def _evaluate_field_condition(condition: str, replacements: dict) -> bool:
    match = re.match(r'^Поле (<[^>]+>) (содержит|равно):\s*(.*)$', condition, re.S)
    if not match:
        return False

    placeholder, operation, values_block = match.groups()
    field_value = _normalize_text(replacements.get(placeholder, ""))
    options = []
    for raw_line in values_block.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        line = re.sub(r'^\d+\.\s*', '', line)
        options.append(_normalize_text(line))

    if not field_value or not options:
        return False
    if operation == "содержит":
        return any(field_value == option or option in field_value for option in options)
    return any(field_value == option for option in options)


def _evaluate_simple_comparison(condition: str, replacements: dict) -> bool:
    equal_match = re.match(r'^Поле (<[^>]+>) = (.+)$', condition)
    if equal_match:
        placeholder, value = equal_match.groups()
        return _normalize_text(replacements.get(placeholder, "")) == _normalize_text(value)

    not_equal_match = re.match(r'^Поле (<[^>]+>) не равно "(.+)"$', condition)
    if not_equal_match:
        placeholder, value = not_equal_match.groups()
        return _normalize_text(replacements.get(placeholder, "")) != _normalize_text(value)

    return False


def _evaluate_age_condition(condition: str, replacements: dict) -> bool:
    normalized_condition = " ".join(
        str(condition or "").replace("\xa0", " ").replace("ё", "е").lower().split()
    )
    if "рабочая дата" not in normalized_condition or "дата рождения" not in normalized_condition:
        return False

    age = _calculate_age_on_date(replacements.get("<Рабочая дата>", ""), replacements.get("<Дата рождения>", ""))
    if age is None:
        return False

    le_match = re.search(r'<=\s*(\d+)', normalized_condition)
    if le_match:
        return age <= int(le_match.group(1))

    lt_match = re.search(r'(?<!<)<\s*(\d+)', normalized_condition)
    if lt_match:
        return age < int(lt_match.group(1))

    return False


def _evaluate_checkbox_condition(condition: str, replacements: dict) -> bool:
    normalized_condition = _normalize_text(condition)

    patent_checked = _is_truthy(replacements.get("<Чекбокс Патент>", ""))
    permit_checked = _is_truthy(replacements.get("<Чекбокс Разрешение>", ""))
    urgent_contract_checked = _is_truthy(replacements.get("<Срочный договор>", "")) or bool(str(replacements.get("<Срочный трудовой договор до>", "") or "").strip())
    ktu_checked = _is_truthy(replacements.get("<В организации применяется КТУ>", ""))
    foreign_citizen_checked = _is_truthy(replacements.get("<Иностранный гражданин>", ""))
    foreign_worker_checked = _is_truthy(replacements.get("<Иностранный работник>", ""))

    if "чек-бокс" in normalized_condition or "чекбокс" in normalized_condition or "чек-бокс" in normalized_condition:
        if "пуст" in normalized_condition and "патент" in normalized_condition and ("разреш" in normalized_condition or "разрещ" in normalized_condition):
            return not patent_checked and not permit_checked
        if "патент" in normalized_condition and "стоит галочка" in normalized_condition:
            return patent_checked
        if ("разрешен" in normalized_condition or "разрещен" in normalized_condition) and "стоит галочка" in normalized_condition:
            return permit_checked

        if "срочн" in normalized_condition and "трудовой договор" in normalized_condition and "стоит галочка" in normalized_condition:
            return urgent_contract_checked
        if "кту" in normalized_condition and "стоит галочка" in normalized_condition:
            return ktu_checked
        if "иностранный гражданин" in normalized_condition and "стоит галочка" in normalized_condition:
            return foreign_citizen_checked
        if "иностранный работник" in normalized_condition and "стоит галочка" in normalized_condition:
            return foreign_worker_checked

    checkbox_map = {
        'В чек-боксе "Срочный трудовой договор до" стоит галочка': "<Срочный трудовой договор до>",
        'В чекбоксе "В организации применяется КТУ" стоит галочка': "<В организации применяется КТУ>",
        'В чек-боксе "Иностранный гражданин" стоит галочка': "<Иностранный гражданин>",
        'В чекбоксе "Иностранный гражданин" стоит галочка': "<Иностранный гражданин>",
        'В чек-боксе "Иностранный работник" стоит галочка': "<Иностранный работник>",
        'В чекбоксе "Иностранный работник" стоит галочка': "<Иностранный работник>",
    }
    placeholder = checkbox_map.get(condition)
    if not placeholder:
        return False

    value = _normalize_text(replacements.get(placeholder, ""))
    return value not in ("", "0", "false", "ложь", "none", "нет", "n")


def _evaluate_section_condition(condition: str, replacements: dict) -> bool:
    match = re.match(r'^Раздел "(.+)" содержит "(.+)"$', condition)
    if not match:
        return False

    section_name, expected = match.groups()
    placeholder = f"<{section_name}>"
    return _normalize_text(expected) in _normalize_text(replacements.get(placeholder, ""))


def _split_compound_condition(condition: str) -> list[str]:
    parts = [part.strip(" ,") for part in re.split(r",\s*а также\s+", condition, flags=re.I) if part.strip(" ,")]
    normalized_parts: list[str] = []
    for part in parts:
        if part.lower().startswith("поле "):
            normalized_parts.append(f"Поле {part[5:]}")
        else:
            normalized_parts.append(part)
    return normalized_parts


def _split_condition_blocks(condition: str) -> list[str]:
    lines = [line.strip() for line in str(condition or "").splitlines() if line.strip()]
    if not lines:
        return []

    block_starts = (
        "Поле ",
        "<Рабочая дата>",
        "В чек-боксе ",
        "В чекбоксе ",
        'Раздел "',
    )

    blocks: list[list[str]] = []
    current_block: list[str] = []

    for line in lines:
        if current_block and any(line.startswith(prefix) for prefix in block_starts):
            blocks.append(current_block)
            current_block = [line]
            continue
        current_block.append(line)

    if current_block:
        blocks.append(current_block)

    return ["\n".join(block).strip() for block in blocks if block]


def _single_condition_matches(condition: str, replacements: dict) -> bool:
    condition = condition.strip()
    if not condition:
        return False

    compound_parts = _split_compound_condition(condition)
    if len(compound_parts) > 1:
        return all(_single_condition_matches(part, replacements) for part in compound_parts)

    return any(
        (
            _evaluate_field_condition(condition, replacements),
            _evaluate_simple_comparison(condition, replacements),
            _evaluate_age_condition(condition, replacements),
            _evaluate_checkbox_condition(condition, replacements),
            _evaluate_section_condition(condition, replacements),
        )
    )

def _rule_matches(condition: str, replacements: dict) -> bool:
    blocks = _split_condition_blocks(condition)
    if not blocks:
        return False
    return all(_single_condition_matches(block, replacements) for block in blocks)


def _rule_matches_with_foreign_fallback(rule: dict, replacements: dict) -> bool:
    point = str(rule.get("point") or "").strip()
    if point == "2.7.":
        urgent_basis = str(replacements.get("<Основание срочного договора>", "") or "").strip()
        if not urgent_basis:
            _write_tkrf_debug("Skipping point 2.7 because urgent contract basis is empty")
            return False
    return _rule_matches(rule["condition"], replacements)


def _debug_urgent_contract_rule(rule: dict, replacements: dict, matched: bool) -> None:
    point = str(rule.get("point") or "").strip()
    condition = str(rule.get("condition") or "").strip()
    if point not in {"2.1.", "2.7."}:
        return
    if "срочн" not in _normalize_text(condition):
        return

    _write_tkrf_debug("=== TKRF URGENT DEBUG START ===")
    _write_tkrf_debug(f"point: {point}")
    _write_tkrf_debug(f"condition: {condition}")
    _write_tkrf_debug(f"matched: {matched}")
    _write_tkrf_debug(f"urgent_until: {replacements.get('<Срочный трудовой договор до>', '')}")
    _write_tkrf_debug(f"urgent_basis: {replacements.get('<Основание срочного договора>', '')}")
    _write_tkrf_debug("=== TKRF URGENT DEBUG END ===")


def _debug_point_1_3_rule(rule: dict, replacements: dict, matched: bool) -> None:
    point = str(rule.get("point") or "").strip()
    if point != "1.3.":
        return

    _write_tkrf_debug("=== TKRF POINT 1.3 DEBUG START ===")
    _write_tkrf_debug(f"rule_index: {rule.get('index')}")
    _write_tkrf_debug(f"condition: {rule.get('condition')}")
    _write_tkrf_debug(f"matched: {matched}")
    _write_tkrf_debug(f"citizenship: {replacements.get('<Гражданство страны>', '')}")
    _write_tkrf_debug(f"foreign_citizen: {replacements.get('<Иностранный гражданин>', '')}")
    _write_tkrf_debug(f"foreign_worker: {replacements.get('<Иностранный работник>', '')}")
    _write_tkrf_debug(f"permit_checkbox: {replacements.get('<Чекбокс Разрешение>', '')}")
    _write_tkrf_debug(f"patent_checkbox: {replacements.get('<Чекбокс Патент>', '')}")
    _write_tkrf_debug(f"document_type: {replacements.get('<Вид документа>', '')}")
    _write_tkrf_debug(f"work_date: {replacements.get('<Рабочая дата>', '')}")
    _write_tkrf_debug(f"birth_date: {replacements.get('<Дата рождения>', '')}")
    _write_tkrf_debug("=== TKRF POINT 1.3 DEBUG END ===")


def _is_russian_citizenship(parsed_data_1c: dict) -> bool:
    country = _normalize_text(parsed_data_1c.get("personal_info", {}).get("Страна", ""))
    return country in {"россия", "российская федерация"}


def _capture_paragraph_template(paragraph: Paragraph) -> dict:
    run_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            run_rpr = deepcopy(run._r.rPr)
            break

    return {
        "style": getattr(paragraph, "style", None),
        "ppr": deepcopy(paragraph._p.pPr) if paragraph._p.pPr is not None else None,
        "run_rpr": run_rpr,
    }


def _apply_paragraph_template(paragraph: Paragraph, template: dict, run=None) -> None:
    style = template.get("style")
    if style is not None:
        try:
            paragraph.style = style
        except Exception:
            pass

    ppr = template.get("ppr")
    if ppr is not None:
        try:
            existing_ppr = paragraph._p.pPr
            if existing_ppr is not None:
                paragraph._p.remove(existing_ppr)
            paragraph._p.insert(0, deepcopy(ppr))
        except Exception:
            pass

    run_rpr = template.get("run_rpr")
    if run is not None and run_rpr is not None:
        try:
            existing_rpr = run._r.rPr
            if existing_rpr is not None:
                run._r.remove(existing_rpr)
            run._r.insert(0, deepcopy(run_rpr))
        except Exception:
            pass


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    template = _capture_paragraph_template(paragraph)
    paragraph.clear()
    run = paragraph.add_run(text)
    _apply_paragraph_template(paragraph, template, run)


def _set_run_format(run, run_rpr) -> None:
    if run_rpr is None:
        return
    try:
        existing_rpr = run._r.rPr
        if existing_rpr is not None:
            run._r.remove(existing_rpr)
        run._r.insert(0, deepcopy(run_rpr))
    except Exception:
        pass


def _replace_placeholders_preserving_runs(paragraph: Paragraph, replacements: dict) -> bool:
    runs = list(paragraph.runs)
    if not runs:
        return False

    full_text = "".join(run.text for run in runs)
    updated_text = _replace_placeholders(full_text, replacements)
    if updated_text == full_text:
        return False

    char_run_indexes: list[int] = []
    run_templates: dict[int, object] = {}
    for run_index, run in enumerate(runs):
        run_templates[run_index] = deepcopy(run._r.rPr) if run._r.rPr is not None else None
        char_run_indexes.extend([run_index] * len(run.text))

    if not char_run_indexes:
        return False

    segments: list[tuple[str, int | None]] = []

    def choose_placeholder_run_index(start: int, end: int) -> int | None:
        previous_index = char_run_indexes[start - 1] if start > 0 else None
        next_index = char_run_indexes[end] if end < len(char_run_indexes) else None

        if next_index is not None:
            return next_index
        if previous_index is not None:
            return previous_index
        if start < len(char_run_indexes):
            return char_run_indexes[start]
        if char_run_indexes:
            return char_run_indexes[-1]
        return None

    def append_literal(start: int, end: int) -> None:
        if end <= start:
            return
        current_run_index = char_run_indexes[start]
        chunk: list[str] = []
        for position in range(start, end):
            run_index = char_run_indexes[position]
            char = full_text[position]
            if run_index != current_run_index:
                segments.append(("".join(chunk), current_run_index))
                chunk = [char]
                current_run_index = run_index
            else:
                chunk.append(char)
        if chunk:
            segments.append(("".join(chunk), current_run_index))

    current_position = 0
    for match in _PLACEHOLDER_PATTERN.finditer(full_text):
        start, end = match.span()
        append_literal(current_position, start)

        placeholder = match.group(0)
        replacement_text = _replace_placeholders(placeholder, replacements)
        source_run_index = choose_placeholder_run_index(start, end)
        if replacement_text:
            segments.append((replacement_text, source_run_index))

        current_position = end

    append_literal(current_position, len(full_text))

    paragraph.clear()
    for segment_text, run_index in segments:
        if not segment_text:
            continue
        new_run = paragraph.add_run(segment_text)
        _set_run_format(new_run, run_templates.get(run_index))

    return True


def _replace_placeholders_in_paragraph(paragraph: Paragraph, replacements: dict) -> None:
    if not paragraph.text:
        return

    updated_runs = False
    for run in paragraph.runs:
        updated_text = _replace_placeholders(run.text, replacements)
        if updated_text != run.text:
            run.text = updated_text
            updated_runs = True

    if updated_runs:
        return

    if _replace_placeholders_preserving_runs(paragraph, replacements):
        return

    updated_text = _replace_placeholders(paragraph.text, replacements)
    if updated_text != paragraph.text:
        _set_paragraph_text(paragraph, updated_text)


def _find_paragraph_by_point(document: Document, point: str) -> Paragraph | None:
    prefix = point.strip()
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    return None


def _find_paragraph_containing_text(document: Document, text: str) -> Paragraph | None:
    search_text = text.strip()
    if not search_text:
        return None

    for paragraph in document.paragraphs:
        if search_text in paragraph.text.strip():
            return paragraph

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if search_text in paragraph.text.strip():
                        return paragraph
    return None


def _find_cell_and_paragraph_containing_text(document: Document, text: str):
    search_text = text.strip()
    if not search_text:
        return None, None

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if search_text in paragraph.text.strip():
                        return cell, paragraph
    return None, None


def _append_text_to_cell(cell, text: str, anchor_paragraph: Paragraph | None = None) -> None:
    text = text.strip()
    if not text:
        return

    for paragraph in cell.paragraphs:
        if text == paragraph.text.strip():
            return

    template = _capture_paragraph_template(anchor_paragraph) if anchor_paragraph is not None else {}
    new_paragraph = cell.add_paragraph()
    run = new_paragraph.add_run(text)
    _apply_paragraph_template(new_paragraph, template, run)


def _insert_paragraph_after(paragraph: Paragraph, text: str, template_paragraph: Paragraph | None = None) -> Paragraph:
    template_source = template_paragraph or paragraph
    template = _capture_paragraph_template(template_source)
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    run = new_paragraph.add_run(text)
    _apply_paragraph_template(new_paragraph, template, run)
    return new_paragraph


def _previous_point(point: str) -> str | None:
    match = re.match(r'^(\d+)\.(\d+)\.$', point.strip())
    if not match:
        return None
    section, number = match.groups()
    previous_number = int(number) - 1
    if previous_number <= 0:
        return None
    return f"{section}.{previous_number}."


def _section_header(point: str) -> str | None:
    match = re.match(r'^(\d+)\.\d+\.$', point.strip())
    if not match:
        return None
    return f"{match.group(1)}."


def _find_insert_anchor(document: Document, point: str) -> Paragraph | None:
    previous_point = _previous_point(point)
    if previous_point:
        paragraph = _find_paragraph_by_point(document, previous_point)
        if paragraph:
            return paragraph

    header_prefix = _section_header(point)
    if header_prefix:
        candidates = [p for p in document.paragraphs if p.text.strip().startswith(header_prefix)]
        if candidates:
            return candidates[-1]

    return None


def _find_template_donor(document: Document, point: str, anchor: Paragraph | None) -> Paragraph | None:
    previous_point = _previous_point(point)
    if previous_point:
        paragraph = _find_paragraph_by_point(document, previous_point)
        if paragraph is not None:
            return paragraph

    header_prefix = _section_header(point)
    if header_prefix:
        section_point_pattern = re.compile(rf"^{re.escape(header_prefix)}\d+\.")
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if section_point_pattern.match(text):
                return paragraph

    return anchor


def _replace_existing_point(document: Document, point: str, text: str) -> None:
    if not re.match(r'^\d+(\.\d+)*\.$', point.strip()):
        paragraph = _find_paragraph_containing_text(document, point)
        if paragraph is not None:
            _set_paragraph_text(paragraph, text)
            return
        document.add_paragraph(text)
        return

    paragraph = _find_paragraph_by_point(document, point)
    if paragraph is None:
        anchor = _find_insert_anchor(document, point)
        if anchor is not None:
            _insert_paragraph_after(anchor, f"{point} {text}")
        return
    _set_paragraph_text(paragraph, f"{point} {text}")


def _add_new_point(document: Document, point: str, text: str) -> None:
    if not re.match(r'^\d+(\.\d+)*\.$', point.strip()):
        cell, anchor = _find_cell_and_paragraph_containing_text(document, point)
        if cell is not None:
            _append_text_to_cell(cell, text, anchor)
            return

        anchor = _find_paragraph_containing_text(document, point)
        if anchor is None:
            document.add_paragraph(text)
            return
        _insert_paragraph_after(anchor, text, anchor)
        return

    if _find_paragraph_by_point(document, point) is not None:
        return
    anchor = _find_insert_anchor(document, point)
    if anchor is None:
        document.add_paragraph(f"{point} {text}")
        return
    template_donor = _find_template_donor(document, point, anchor)
    _insert_paragraph_after(anchor, f"{point} {text}", template_donor)


def _append_to_existing_point(document: Document, point: str, text: str) -> None:
    if not re.match(r'^\d+(\.\d+)*\.$', point.strip()):
        cell, paragraph = _find_cell_and_paragraph_containing_text(document, point)
        if cell is not None:
            _append_text_to_cell(cell, text, paragraph)
            return

        paragraph = _find_paragraph_containing_text(document, point)
        if paragraph is None:
            _add_new_point(document, point, text)
            return
        existing_text = paragraph.text.strip()
        if text in existing_text:
            return
        separator = "\n" if existing_text else ""
        _set_paragraph_text(paragraph, f"{existing_text}{separator}{text}")
        return

    paragraph = _find_paragraph_by_point(document, point)
    if paragraph is None:
        _add_new_point(document, point, text)
        return
    existing_text = paragraph.text.strip()
    if text in existing_text:
        return
    separator = "\n" if existing_text else ""
    _set_paragraph_text(paragraph, f"{existing_text}{separator}{text}")


def _apply_formulation_rules(document: Document, parsed_data_1c: dict) -> None:
    replacements = _build_contract_placeholders(parsed_data_1c)
    rules = _load_formulation_rules()
    russian_citizenship = _is_russian_citizenship(parsed_data_1c)
    job_info = parsed_data_1c.get("job_info", {})
    personal_info = parsed_data_1c.get("personal_info", {})
    document_info = parsed_data_1c.get("document_info", {})
    hire_date_raw = job_info.get("ДатаПриемаНаРаботу", "")
    document_birth_date_raw = document_info.get("ДатаРождения", "")
    personal_birth_date_raw = personal_info.get("ДатаРождения", "")
    resolved_birth_date = replacements.get("<Дата рождения>", "")
    parsed_hire_date = _parse_date(hire_date_raw)
    parsed_birth_date = _parse_date(resolved_birth_date)
    employee_age = _calculate_age_on_date(
        replacements.get("<Рабочая дата>", ""),
        replacements.get("<Дата рождения>", ""),
    )
    _write_minor_contract_debug("=== MINOR CONTRACT DEBUG START ===")
    _write_minor_contract_debug(
        "minor_contract_input: "
        f"fio={personal_info.get('ФИО', '')!r}; "
        f"hire_date_job_info={hire_date_raw!r}; "
        f"birth_date_document_info={document_birth_date_raw!r}; "
        f"birth_date_personal_info={personal_birth_date_raw!r}; "
        f"birth_date_selected={resolved_birth_date!r}; "
        f"hire_date_parsed={parsed_hire_date!r}; "
        f"birth_date_parsed={parsed_birth_date!r}; "
        f"calculated_age={employee_age!r}; "
        f"is_minor={employee_age is not None and employee_age < 18}; "
        f"russian_citizenship={russian_citizenship}"
    )
    skip_points = {
        "ИнформацияОПатенте",
        'Реквизиты Работника в разделе "Адреса и реквизиты Сторон" трудового договора',
    }

    matched_age_rules: list[str] = []
    for rule in rules:
        point = str(rule["point"] or "").strip()
        if point in skip_points:
            continue
        is_age_based_rule = "дата рождения" in _normalize_text(rule["condition"])
        if point == "1.3." and russian_citizenship and not is_age_based_rule:
            _write_tkrf_debug("Skipping point 1.3 for Russian citizenship")
            continue
        matched = _rule_matches_with_foreign_fallback(rule, replacements)
        if is_age_based_rule:
            action = str(rule["action"] or "").strip()
            _write_minor_contract_debug(
                "minor_contract_rule: "
                f"row={rule.get('index')!r}; point={point!r}; action={action!r}; "
                f"condition={rule['condition']!r}; "
                f"age_condition_matches={_evaluate_age_condition(rule['condition'], replacements)}; "
                f"rule_matches={matched}"
            )
            if matched:
                matched_age_rules.append(f"{point} ({action})")
        _debug_urgent_contract_rule(rule, replacements, matched)
        _debug_point_1_3_rule(rule, replacements, matched)
        if not matched:
            continue

        rendered_text = _replace_placeholders(rule["text"], replacements)
        action = str(rule["action"] or "").strip()

        if action == "Добавить пункт":
            _add_new_point(document, point, rendered_text)
        elif action == "Заменить существующий пункт":
            _replace_existing_point(document, point, rendered_text)
        elif action == "Добавить в существующий пункт":
            _append_to_existing_point(document, point, rendered_text)

    _write_minor_contract_debug(
        "minor_contract_result: "
        f"calculated_age={employee_age!r}; matched_age_rules={matched_age_rules!r}"
    )
    _write_minor_contract_debug("=== MINOR CONTRACT DEBUG END ===")

    if _is_hourly_payment(parsed_data_1c, replacements):
        hourly_text = _get_hourly_payment_rule_text(rules, replacements)
        _replace_existing_point(document, "6.1.", hourly_text)
        _write_tkrf_debug(
            "hourly_payment_clause_applied: "
            f"hourly_rate={replacements.get('<Оплата по часовому тарифу>', '')!r}; "
            f"accrual={replacements.get('<Начисление>', '')!r}"
        )


def _replace_placeholders_in_paragraphs(document: Document, replacements: dict) -> None:
    for paragraph in document.paragraphs:
        _replace_placeholders_in_paragraph(paragraph, replacements)


def _replace_placeholders_in_tables(document: Document, replacements: dict) -> None:
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_placeholders_in_paragraph(paragraph, replacements)


def _replace_placeholders_in_story(story, replacements: dict) -> None:
    for paragraph in story.paragraphs:
        _replace_placeholders_in_paragraph(paragraph, replacements)

    for table in story.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_placeholders_in_paragraph(paragraph, replacements)


def _find_run_ancestor(element, parent_map):
    current = element
    while current is not None:
        if current.tag == f"{_W_TAG}r":
            return current
        current = parent_map.get(current)
    return None


def _copy_run_properties(source_run, target_run) -> None:
    if source_run is None or target_run is None:
        return

    source_rpr = source_run.find(f"{_W_TAG}rPr")
    target_rpr = target_run.find(f"{_W_TAG}rPr")

    if target_rpr is not None:
        target_run.remove(target_rpr)
    if source_rpr is not None:
        target_run.insert(0, deepcopy(source_rpr))


def _replace_fragmented_placeholders_in_xml_part(xml_bytes: bytes, replacements: dict) -> bytes:
    try:
        root = ET.fromstring(xml_bytes)
    except ET.ParseError:
        return xml_bytes

    text_nodes = [node for node in root.iter(f"{_W_TAG}t")]
    if not text_nodes:
        return xml_bytes

    normalized_replacements = {
        _normalize_placeholder_name(placeholder): str(value)
        for placeholder, value in replacements.items()
    }
    parent_map = {child: parent for parent in root.iter() for child in parent}

    changed = False
    index = 0
    while index < len(text_nodes):
        node = text_nodes[index]
        node_text = node.text or ""

        if "<" not in node_text and node_text != "<":
            index += 1
            continue

        candidate_text = node_text
        end_index = index
        while ">" not in candidate_text and end_index + 1 < len(text_nodes):
            end_index += 1
            candidate_text += text_nodes[end_index].text or ""

        stripped_candidate = candidate_text.strip()
        if not (stripped_candidate.startswith("<") and stripped_candidate.endswith(">")):
            index += 1
            continue

        replacement_value = normalized_replacements.get(_normalize_placeholder_name(stripped_candidate))
        if replacement_value is None:
            index += 1
            continue

        donor_node = text_nodes[end_index + 1] if end_index + 1 < len(text_nodes) else (text_nodes[index - 1] if index > 0 else None)
        target_run = _find_run_ancestor(text_nodes[index], parent_map)
        donor_run = _find_run_ancestor(donor_node, parent_map) if donor_node is not None else None
        if donor_run is not None and target_run is not None:
            _copy_run_properties(donor_run, target_run)

        text_nodes[index].text = replacement_value
        for clear_index in range(index + 1, end_index + 1):
            text_nodes[clear_index].text = ""

        changed = True
        index = end_index + 1

    if not changed:
        return xml_bytes

    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _postprocess_fragmented_placeholders_in_docx(output_path: str, replacements: dict) -> None:
    docx_path = Path(output_path)
    if not docx_path.exists():
        return

    target_parts_pattern = re.compile(r"^word/(document|header\d+|footer\d+)\.xml$")

    with zipfile.ZipFile(docx_path, "r") as source_zip:
        file_map = {name: source_zip.read(name) for name in source_zip.namelist()}

    changed = False
    for name, content in list(file_map.items()):
        if not target_parts_pattern.match(name):
            continue
        updated_content = _replace_fragmented_placeholders_in_xml_part(content, replacements)
        if updated_content != content:
            file_map[name] = updated_content
            changed = True

    if not changed:
        return

    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as target_zip:
        for name, content in file_map.items():
            target_zip.writestr(name, content)


def _repair_docx_with_word(output_path: str) -> None:
    if win32com is None:
        _write_tkrf_debug("Word repair skipped: win32com is unavailable")
        return

    word_app = None
    document = None
    try:
        word_app = win32com.client.DispatchEx("Word.Application")
        word_app.Visible = False
        word_app.DisplayAlerts = 0

        document = word_app.Documents.Open(
            FileName=str(Path(output_path).resolve()),
            ConfirmConversions=False,
            ReadOnly=False,
            AddToRecentFiles=False,
            Visible=False,
            OpenAndRepair=True,
        )
        document.Save()
        _write_tkrf_debug("Word repair completed successfully")
    except Exception as error:
        _write_tkrf_debug(f"Word repair failed: {error}")
    finally:
        if document is not None:
            try:
                document.Close(SaveChanges=False)
            except Exception:
                pass
        if word_app is not None:
            try:
                word_app.Quit()
            except Exception:
                pass


def tkrf_sample_former(input_sample: str, output_result: str, parsed_data_1c: dict) -> str:
    replacements = _build_contract_placeholders(parsed_data_1c)
    replacements = _apply_rule_driven_placeholders(parsed_data_1c, replacements)
    replacements = {key.strip("<>"): value for key, value in replacements.items()}
    _write_tkrf_debug("=== TKRF GENERATION START ===")
    _write_tkrf_debug(f"input_sample: {input_sample}")
    _write_tkrf_debug(f"output_result: {output_result}")
    _write_tkrf_debug(f"urgent_until_replacement: {replacements.get('Срочный трудовой договор до', '')}")
    _write_tkrf_debug(f"urgent_basis_replacement: {replacements.get('Основание срочного договора', '')}")


    try:
        document = Document(input_sample)
        placeholder_replacements = {f"<{k}>": v for k, v in replacements.items()}
        _replace_placeholders_in_story(document, placeholder_replacements)
        for section in document.sections:
            _replace_placeholders_in_story(section.header, placeholder_replacements)
            _replace_placeholders_in_story(section.footer, placeholder_replacements)
        _apply_formulation_rules(document, parsed_data_1c)
        _replace_placeholders_in_story(document, placeholder_replacements)
        for section in document.sections:
            _replace_placeholders_in_story(section.header, placeholder_replacements)
            _replace_placeholders_in_story(section.footer, placeholder_replacements)
        document.save(output_result)
        _repair_docx_with_word(output_result)
        _write_tkrf_debug("=== TKRF GENERATION END ===")
        return output_result
    except Exception as e:
        _write_tkrf_debug(f"TKRF generation error: {e}")
        return f"Ошибка во время формирования файла: {e}"
