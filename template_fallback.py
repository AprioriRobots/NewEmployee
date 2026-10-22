import re
from datetime import datetime
from typing import Dict, Iterable, Optional

from docx import Document
from openpyxl import load_workbook

from name_utils import get_fio_part, get_name_patronymic_initials
from template_utils import expand_relative_month_keys


_PLACEHOLDER_RE = re.compile(r"<([^<>]+)>")


def _to_int(value) -> Optional[int]:
    if value in (None, ""):
        return None
    try:
        if isinstance(value, str):
            cleaned = value.replace(" ", "").replace(",", ".")
            return int(float(cleaned))
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _choose_plural(value: int, forms) -> str:
    value = abs(value) % 100
    if 11 <= value <= 19:
        return forms[2]
    last_digit = value % 10
    if last_digit == 1:
        return forms[0]
    if 2 <= last_digit <= 4:
        return forms[1]
    return forms[2]


def number_to_words_ru(value) -> str:
    number = _to_int(value)
    if number is None:
        return ""
    if number == 0:
        return "ноль"

    ones_male = ("", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять")
    ones_female = ("", "одна", "две", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять")
    teens = (
        "десять", "одиннадцать", "двенадцать", "тринадцать", "четырнадцать",
        "пятнадцать", "шестнадцать", "семнадцать", "восемнадцать", "девятнадцать",
    )
    tens = ("", "", "двадцать", "тридцать", "сорок", "пятьдесят", "шестьдесят", "семьдесят", "восемьдесят", "девяносто")
    hundreds = ("", "сто", "двести", "триста", "четыреста", "пятьсот", "шестьсот", "семьсот", "восемьсот", "девятьсот")
    orders = (
        (("", "", ""), False),
        (("тысяча", "тысячи", "тысяч"), True),
        (("миллион", "миллиона", "миллионов"), False),
        (("миллиард", "миллиарда", "миллиардов"), False),
    )

    def triplet_to_words(triplet: int, female: bool):
        parts = []
        parts.append(hundreds[triplet // 100])
        remainder = triplet % 100
        if 10 <= remainder <= 19:
            parts.append(teens[remainder - 10])
        else:
            parts.append(tens[remainder // 10])
            ones = ones_female if female else ones_male
            parts.append(ones[remainder % 10])
        return [part for part in parts if part]

    parts = []
    order_index = 0
    while number > 0 and order_index < len(orders):
        triplet = number % 1000
        if triplet:
            order_forms, is_female = orders[order_index]
            triplet_words = triplet_to_words(triplet, is_female)
            if order_index > 0:
                triplet_words.append(_choose_plural(triplet, order_forms))
            parts.insert(0, " ".join(triplet_words))
        number //= 1000
        order_index += 1

    return " ".join(parts).strip()


def build_common_replacements(parsed_data_1c: dict) -> Dict[str, str]:
    job_info = parsed_data_1c.get("job_info", {})
    org_info = job_info.get("Организация", {})
    personal_info = parsed_data_1c.get("personal_info", {})
    document_info = parsed_data_1c.get("document_info", {})
    military_info = parsed_data_1c.get("military_info", {})
    fio = personal_info.get("ФИО", "")

    replacements = {
        "Фамилия": get_fio_part(fio, 0) if fio else "",
        "Имя": get_fio_part(fio, 1) if fio else "",
        "Отчество": get_fio_part(fio, 2) if fio else "",
        "Полное имя работника": fio,
        "Инициалы имени, отчества": get_name_patronymic_initials(fio) if fio else "",
        "Пол": personal_info.get("ПОЛ", ""),
        "Пол работника": personal_info.get("ПОЛ", ""),
        "Подразделение": job_info.get("Подразделение", ""),
        "Должность работника": job_info.get("Должность", ""),
        "Табельный номер работника": job_info.get("ТабельныйНомерРаботника", job_info.get("ТабельныйНомер", "")),
        "Дата приема": job_info.get("ДатаПриемаНаРаботу", ""),
        "Дата приказа о приеме": job_info.get("ДатаПриемаНаРаботу", ""),
        "Испытательный срок": str(job_info.get("ИспытательныйСрок", "")),
        "Испытательный срок словами": number_to_words_ru(job_info.get("ИспытательныйСрок", "")),
        "График работы": job_info.get("ГрафикРаботы", ""),
        "Оклад": str(job_info.get("Оклад", "")),
        "Оклад словами": number_to_words_ru(job_info.get("Оклад", "")),
        "Количество дней ежегодного отпуска": str(job_info.get("Отпуск", "")),
        "Длительность рабочей недели": str(job_info.get("ДлительностьРабочейНедели", "")),
        "Договор №": job_info.get("ТкНомер", ""),
        "Номер приказа о приеме": job_info.get("НомерПриказаОПриеме", job_info.get("ТкНомер", "")),
        "Приказ №": job_info.get("НомерПриказаОПриеме", job_info.get("ТкНомер", "")),
        "Основание подписи": job_info.get("ОснованиеПодписи", ""),
        "Основания подписи": job_info.get("ОснованиеПодписи", ""),
        "Полное наименование": org_info.get("НаименованиеПолное", ""),
        "Краткое наименование": org_info.get("НаименованиеСокращенное", ""),
        "Должность руководителя": org_info.get("ДолжностьРуководителя", ""),
        "Руководитель": org_info.get("Руководитель", ""),
        "Юридический адрес": org_info.get("ЮридическийАдрес", ""),
        "Фактический адрес": org_info.get("ФактическийАдрес", ""),
        "ИНН организации": org_info.get("ИНН_Орг", ""),
        "КПП организации": org_info.get("КПП_Орг", ""),
        "ОГРН организации": org_info.get("ОГРН(ИП)", ""),
        "Расчетный счет организации": org_info.get("РасСчет", ""),
        "БИК банка организации": org_info.get("БикБанка", ""),
        "Корреспондентский счет организации": org_info.get("КорСчет", ""),
        "Дата рождения": document_info.get("ДатаРождения", ""),
        "Место рождения": document_info.get("МестоРождения", ""),
        "Серия документа, удостоверяющего личность": document_info.get("Серия", ""),
        "Номер документа, удостоверяющего личность": document_info.get("Номер", ""),
        "Дата выдачи документа, удостоверяющего личность": document_info.get("ДатаВыдачи", ""),
        "Орган, выдавший документ, удостоверяющего личность": document_info.get("КемВыдан", ""),
        "Код подразделения": document_info.get("КодПодразделения", ""),
        "СНИЛС": personal_info.get("СНИЛС", ""),
        "Адрес по прописке": personal_info.get("Адрес_прописка", ""),
        "Адрес места проживания": personal_info.get("Адрес_проживания", ""),
        "Мобильный телефон": personal_info.get("Телефон", ""),
        "Банк работника": personal_info.get("Банк", ""),
        "Адрес банка работника": personal_info.get("АдресБанка", ""),
        "ИНН банка работника": personal_info.get("ИннБанка", ""),
        "БИК банка работника": personal_info.get("БикБанка", ""),
        "Корреспондентский счёт банка работника": personal_info.get("КоррСчет", ""),
        "Номер счёта работника": personal_info.get("НомерСчета", ""),
        "Военкомат": military_info.get("Военкомат", ""),
        "Воинское звание": military_info.get("ВоинскоеЗвание", ""),
        "Категория запаса": military_info.get("КатегорияЗапаса", ""),
        "Состав (профиль)": military_info.get("СоставПрофиль", ""),
        "ВУС": military_info.get("ВУС", ""),
        "Годность к воинской службе": military_info.get("Годность", ""),
        "Адрес военкомата": military_info.get("АдресВоенкомата", ""),
        "Документ воинского учёта": military_info.get("ДокументВоинскогоУчета", ""),
        "Серия документа воинского учёта": military_info.get("СерияДокументаВоинскогоУчета", ""),
        "Номер документа воинского учёта": military_info.get("НомерДокументаВоинскогоУчета", ""),
        "Отдел, выдавший документ воинского учёта": military_info.get("ОтделВыдавшийДокументВоинскогоУчета", ""),
        "Дата выдачи документа воинского учёта": military_info.get("ДатаВыдачиДокументаВоинскогоУчета", ""),
        "Должность ответственного за ВУ": military_info.get("ДолжностьОтветственногоЗаВУ", ""),
        "Инициалы имени, отчества ответственного за ВУ": military_info.get("ИнициалыИмениОтчестваОтветственногоЗаВУ", ""),
        "Фамилия ответственного за ВУ": military_info.get("ФамилияОтветсвтенногоЗаВУ", ""),
        "Семейное положение": personal_info.get("СемейноеПоложение", ""),
        "Дата регистрации": personal_info.get("ДатаРегистрации", ""),
        "Вид образования": personal_info.get("Образование", {}).get("ВидОбразования", ""),
        "Учебное заведение": personal_info.get("Образование", {}).get("УчебноеЗаведение", ""),
        "Квалификация": personal_info.get("Образование", {}).get("Квалификация", ""),
        "Специальность": personal_info.get("Образование", {}).get("Специальность", ""),
        "Период обучения по": personal_info.get("Образование", {}).get("ПериодОбучения", ""),
        "Серия квалификационного документа": personal_info.get("Образование", {}).get("Серия", ""),
        "Номер квалификационного документа": personal_info.get("Образование", {}).get("Номер", ""),
        "Должность руководителя": org_info.get("ДолжностьРуководителя", ""),
        "Инициалы имени, отчества руководителя": get_name_patronymic_initials(org_info.get("Руководитель", "")),
        "Фамилия руководителя": get_fio_part(org_info.get("Руководитель", ""), 0),
    }

    relatives = personal_info.get("РодственникиСписок") or []
    if not relatives:
        single_relative = personal_info.get("Родственники", {})
        if isinstance(single_relative, dict):
            relatives = [single_relative]

    relative_fields = (
        ("СтепеньРодства", "Степень родства"),
        ("Фамилия", "Фамилия родственника"),
        ("Имя", "Имя родственника"),
        ("Отчество", "Отчество родственника"),
        ("ДатаРождения", "Дата рождения родственника"),
    )
    for index in range(5):
        relative = relatives[index] if index < len(relatives) else {}
        suffix = "" if index == 0 else str(index + 1)
        for source_key, placeholder_name in relative_fields:
            replacements[f"{placeholder_name}{suffix}"] = relative.get(source_key, "")

    commission_fields = (
        ("ДолжностьПредседателяКомиссии", "Должность председателя комиссии"),
        ("ФамилияПредседателяКомиссии", "Фамилия председателя комиссии"),
        ("ИнициалыИмениОтчестваПредседателяКомиссии", "Инициалы имени, отчества председателя комиссии"),
        ("ДолжностьЧленаКомиссии1", "Должность члена комиссии 1"),
        ("ФамилияЧленаКомиссии1", "Фамилия члена комиссии 1"),
        ("ИнициалыИмениОтчестваЧленаКомиссии1", "Инициалы имени, отчества члена комиссии 1"),
        ("ДолжностьЧленаКомиссии2", "Должность члена комиссии 2"),
        ("ФамилияЧленаКомиссии2", "Фамилия члена комиссии 2"),
        ("ИнициалыИмениОтчестваЧленаКомиссии2", "Инициалы имени, отчества члена комиссии 2"),
    )
    for source_key, placeholder_name in commission_fields:
        replacements[placeholder_name] = org_info.get(source_key, "")

    responsible_mappings = {
        "Фамилия ответственного за ОТ": ("ФамилияОтветственногоЗаОТ", "ФамилияОтветственногоЗаПроведениеИнструктажейПоОТ"),
        "Инициалы имени, отчества ответственного за ОТ": ("ИнициалыИмениОтчестваОтветственногоЗаОТ", "ИнициалыИмениОтчестваОтветственногоЗаПроведениеИнструктажейПоОТ"),
        "Должность ответственного за ОТ": ("ДолжностьОтветственногоЗаОТ", "ДолжностьОтветственногоЗаПроведениеИнструктажейПоОТ"),
        "Фамилия ответственного за проведение инструктажей по ОТ": ("ФамилияОтветственногоЗаПроведениеИнструктажейПоОТ", "ФамилияОтветственногоЗаОТ"),
        "Инициалы имени, отчества ответственного за проведение инструктажей по ОТ": ("ИнициалыИмениОтчестваОтветственногоЗаПроведениеИнструктажейПоОТ", "ИнициалыИмениОтчестваОтветственногоЗаОТ"),
        "Должность ответственного за проведение инструктажей по ОТ": ("ДолжностьОтветственногоЗаПроведениеИнструктажейПоОТ", "ДолжностьОтветственногоЗаОТ"),
        "Фамилия ответственного за ПБ": ("ФамилияОтветственногоЗаПБ", "ФамилияОтветственногоЗаПроведениеИнструктажейПоПБ"),
        "Инициалы имени, отчества ответственного за ПБ": ("ИнициалыИмениОтчестваОтветственногоЗаПБ", "ИнициалыИмениОтчестваОтветственногоЗаПроведениеИнструктажейПоПБ"),
        "Должность ответственного за ПБ": ("ДолжностьОтветственногоЗаПБ", "ДолжностьОтветственногоЗаПроведениеИнструктажейПоПБ"),
        "Фамилия ответственного за проведение инструктажей по ПБ": ("ФамилияОтветственногоЗаПроведениеИнструктажейПоПБ", "ФамилияОтветственногоЗаПБ"),
        "Инициалы имени, отчества ответственного за проведение инструктажей по ПБ": ("ИнициалыИмениОтчестваОтветственногоЗаПроведениеИнструктажейПоПБ", "ИнициалыИмениОтчестваОтветственногоЗаПБ"),
        "Должность ответственного за проведение инструктажей по ПБ": ("ДолжностьОтветственногоЗаПроведениеИнструктажейПоПБ", "ДолжностьОтветственногоЗаПБ"),
        "Фамилия ответственного за выдачу СИЗ": ("ФамилияОтветственногоЗаВыдачуСИЗ", "ФамилияОтветственногоЗаОТ"),
        "Инициалы имени, отчества ответственного за выдачу СИЗ": ("ИнициалыИмениОтчестваОтветственногоЗаВыдачуСИЗ", "ИнициалыИмениОтчестваОтветственногоЗаОТ"),
        "Должность ответственного за выдачу СИЗ": ("ДолжностьОтветственногоЗаВыдачуСИЗ", "ДолжностьОтветственногоЗаОТ"),
        "Фамилия ответственного за ГО и ЧС": ("ФамилияОтветственногоЗаГОиЧС",),
        "Инициалы имени, отчества ответственного за ГО и ЧС": ("ИнициалыИмениОтчестваОтветственногоЗаГОиЧС",),
        "Должность ответственного за ГО и ЧС": ("ДолжностьОтветственногоЗаГОиЧС",),
        "Фамилия ответственного за ЭБ": ("ФамилияОтветственногоЗаЭБ",),
        "Инициалы имени, отчества ответственного за ЭБ": ("ИнициалыИмениОтчестваОтветственногоЗаЭБ",),
        "Должность ответственного за ЭБ": ("ДолжностьОтветственногоЗаЭБ",),
        "Фамилия ответственного за АТЗ": ("ФамилияОтветственногоЗаАТЗ",),
        "Инициалы имени, отчества ответственного за АТЗ": ("ИнициалыИмениОтчестваОтветственногоЗаАТЗ",),
        "Должность ответственного за АТЗ": ("ДолжностьОтветственногоЗаАТЗ",),
    }

    for placeholder_name, org_keys in responsible_mappings.items():
        value = ""
        for org_key in org_keys:
            value = org_info.get(org_key, "")
            if value:
                break
        replacements[placeholder_name] = value

    return replacements


def _replace_text(text: str, replacements: Dict[str, str]) -> str:
    if not text or "<" not in text:
        return text

    candidate_keys = [match.group(1) for match in _PLACEHOLDER_RE.finditer(text)]
    expanded = expand_relative_month_keys(replacements, candidate_keys)
    result = text
    for placeholder_name in candidate_keys:
        replacement = expanded.get(placeholder_name)
        if replacement is None:
            continue
        result = result.replace(f"<{placeholder_name}>", str(replacement))
    return result


def _replace_in_paragraph(paragraph, replacements: Dict[str, str]) -> None:
    if not paragraph.text or "<" not in paragraph.text:
        return

    updated = _replace_text(paragraph.text, replacements)
    if updated == paragraph.text:
        return

    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = updated


def _iter_docx_paragraphs(document: Document) -> Iterable:
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            yield paragraph
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        yield paragraph
        for paragraph in section.footer.paragraphs:
            yield paragraph
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        yield paragraph


def apply_common_placeholder_fallback(file_path: str, parsed_data_1c: dict) -> str:
    replacements = build_common_replacements(parsed_data_1c)

    if file_path.lower().endswith(".docx"):
        document = Document(file_path)
        for paragraph in _iter_docx_paragraphs(document):
            _replace_in_paragraph(paragraph, replacements)
        document.save(file_path)
        return file_path

    if file_path.lower().endswith(".xlsx"):
        workbook = load_workbook(file_path)
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and "<" in cell.value:
                        cell.value = _replace_text(cell.value, replacements)
        workbook.save(file_path)
        return file_path

    return file_path
