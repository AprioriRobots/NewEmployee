from docx import Document
from python_docx_replace import docx_replace, docx_get_keys

from template_utils import expand_relative_month_keys


# def zp_align_sample_former(input_sample: str, output_result: str, parsed_data_1c: dict) -> str:
#     ZP_ALIGN_INPUT_WORDS: dict = {
#         'Должность руководителя': parsed_data_1c['job_info']['Организация']['ДолжностьРуководителя'],
#         'Краткое наименование': parsed_data_1c['job_info']['Организация']['НаименованиеСокращенное'],
#         'Руководитель': parsed_data_1c['job_info']['Организация']['Руководитель'],
#         'Должность работника': parsed_data_1c['job_info']['Должность'],
#         'Полное имя работника': parsed_data_1c['personal_info']['ФИО'],
#         'Банк работника': parsed_data_1c['personal_info']['Банк'],
#         'Корреспондентский счёт банка работника': parsed_data_1c['personal_info']['КоррСчет'],
#         'БИК банка работника': parsed_data_1c['personal_info']['БикБанка'],
#         'ИНН банка работника': parsed_data_1c['personal_info']['ИннБанка'],
#         'Номер счёта работника': parsed_data_1c['personal_info']['НомерСчета']
#     }
#
#     try:
#         doc = Document(input_sample)
#
#         paragraph_count: int = len(doc.paragraphs)
#         table_count: int = len(doc.tables)
#
#         for i in range(table_count):
#             # Итерируемся по каждой строке
#             for row in doc.tables[i].table.rows:
#                 current_row = row
#                 cell_count_of_current_row: int = len(current_row.cells)
#
#                 # Итерируемся по каждой ячейке
#                 for j in range(cell_count_of_current_row):
#                     current_cell_txt = current_row.cells[j].text
#
#                     if '<Должность руководителя>' in current_cell_txt:
#                         update_text = current_cell_txt.replace('<Должность руководителя>',
#                                                                ZP_ALIGN_INPUT_WORDS['Должность_руководителя'])
#                         current_row.cells[j].text = update_text
#                         current_cell_txt = update_text
#
#                     if '<Краткое наименование>' in current_cell_txt:
#                         update_text = current_cell_txt.replace('<Краткое наименование>',
#                                                                ZP_ALIGN_INPUT_WORDS['Краткое_наименование'])
#                         current_row.cells[j].text = update_text
#                         current_cell_txt = update_text
#
#                     if '<Руководитель>' in current_cell_txt:
#                         update_text = current_cell_txt.replace('<Руководитель>',
#                                                                ZP_ALIGN_INPUT_WORDS['Руководитель'])
#                         current_row.cells[j].text = update_text
#                         current_cell_txt = update_text
#
#         for i in range(paragraph_count):
#             print(doc.paragraphs[i].text)
#             current_txt = doc.paragraphs[i].text
#
#             if '<Должность работника>' in current_txt:
#                 updated_text = current_txt.replace('<Должность работника>',
#                                                    ZP_ALIGN_INPUT_WORDS['Должность_работника'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Полное имя работника>' in current_txt:
#                 updated_text = current_txt.replace('<Полное имя работника>',
#                                                    ZP_ALIGN_INPUT_WORDS['Полное_имя_работника'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Банк работника>' in current_txt:
#                 updated_text = current_txt.replace('<Банк работника>',
#                                                    ZP_ALIGN_INPUT_WORDS['Банк_работника'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Корреспондентский счёт банка работника>' in current_txt:
#                 updated_text = current_txt.replace('<Корреспондентский счёт банка работника>',
#                                                    ZP_ALIGN_INPUT_WORDS['Корреспондентский_счёт_банка_работника'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<БИК банка работника>' in current_txt:
#                 updated_text = current_txt.replace('<БИК банка работника>',
#                                                    ZP_ALIGN_INPUT_WORDS['БИК_банка_работника'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<ИНН банка работника>' in current_txt:
#                 updated_text = current_txt.replace('<ИНН банка работника>',
#                                                    ZP_ALIGN_INPUT_WORDS['ИНН_банка_работника'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#             if '<Номер счёта работника>' in current_txt:
#                 updated_text = current_txt.replace('<Номер счёта работника>',
#                                                    ZP_ALIGN_INPUT_WORDS['Номер_счёта_работника'])
#                 doc.paragraphs[i].text = updated_text
#                 current_txt = updated_text
#
#         doc.save(output_result)
#         return f'Файл успешно сформирован. Сохранен в {output_result}'
#     except Exception as e:
#         return f'Ошибка во время формирования файла: {e}'


def zp_align_sample_former(input_sample: str, output_result: str, parsed_data_1c: dict) -> str:
    zp_align_input_words: dict = {
        'Должность руководителя': parsed_data_1c['job_info']['Организация']['ДолжностьРуководителя'],
        'Краткое наименование': parsed_data_1c['job_info']['Организация']['НаименованиеСокращенное'],
        'Руководитель': parsed_data_1c['job_info']['Организация']['Руководитель'],
        'Должность работника': parsed_data_1c['job_info']['Должность'],
        'Полное имя работника': parsed_data_1c['personal_info']['ФИО'],
        'Банк работника': parsed_data_1c['personal_info']['Банк'],

        'Корреспондентский счёт банка работника': parsed_data_1c['personal_info']['КоррСчет'],

        'БИК банка работника': parsed_data_1c['personal_info']['БикБанка'],
        'ИНН банка работника': parsed_data_1c['personal_info']['ИннБанка'],
        'Номер счёта работника': parsed_data_1c['personal_info']['НомерСчета']
    }

    try:
        doc = Document(input_sample)
        keys = docx_get_keys(doc, pattern='<>')
        print(keys)
        docx_replace(doc, '<>', **expand_relative_month_keys(zp_align_input_words, keys))

        doc.save(output_result)
        return output_result

    except Exception as e:
        return f'Ошибка во время формирования файла: {e}'


if __name__ == '__main__':
    result: str = zp_align_sample_former(
        input_sample='input_sample/Заявление на перечисление заработной платы.docx',
        output_result='output_result/output_Заявление на перечисление заработной платы.docx'
    )

