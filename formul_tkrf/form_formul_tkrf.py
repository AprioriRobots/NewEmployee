from openpyxl import load_workbook


def build_formul_tkrf_input_words(parsed_data_1c: dict) -> dict:
    return {
        "<Гражданство страны>": parsed_data_1c["personal_info"]["Страна"],
        "<Дата рождения>": parsed_data_1c["document_info"]["ДатаРождения"],
        "<Длительность рабочей недели>": parsed_data_1c["job_info"]["ГрафикРаботы"],
        "<Количество дней ежегодного отпуска>": "28",
        "<Испытательный срок>": str(parsed_data_1c["job_info"]["ИспытательныйСрок"]),
        "<Должность руководителя>": parsed_data_1c["job_info"]["Организация"]["ДолжностьРуководителя"],
    }


def formul_tkrf_sample_former(input_sample: str, output_result: str, parsed_data_1c: dict) -> str:
    formul_tkrf_input_words = build_formul_tkrf_input_words(parsed_data_1c)

    try:
        wb = load_workbook(input_sample)
        ws = wb["Лист1"]

        for row_index in range(1, 25):
            for column in ("D", "E"):
                cell = ws[f"{column}{row_index}"]
                if not isinstance(cell.value, str):
                    continue

                updated_value = cell.value
                for placeholder, replacement in formul_tkrf_input_words.items():
                    updated_value = updated_value.replace(placeholder, replacement)

                cell.value = updated_value

        wb.save(output_result)
        return output_result
    except Exception as e:
        return f"Ошибка во время формирования файла: {e}"

