from docx import Document
from python_docx_replace import docx_replace, docx_get_keys

from template_utils import expand_relative_month_keys


# Пример извлеченных данных (реальный, взятый из базы ЗУП)
# parsed_data_1c = {'document_info': {'ВидДокумента': 'Паспорт гражданина РФ', 'ДатаВыдачи': '08.11.2017', 'ДатаРождения': '22.06.2003', 'КемВыдан': 'ТП №6 Межрайонного ОУФМС России по Московской обл. в городском поселении Люберцы', 'КодПодразделения': '500-078', 'МестоРождения': 'Москва', 'Номер': '738100', 'Серия': '46 17', 'СрокДокумента': '01.01.0100'}, 'job_info': {'АвансПериод': '', 'ГрафикРаботы': 'Сменный 1 ', 'ДатаПриемаНаРаботу': '26.07.2023', 'Должность': 'Бармен', 'ИспытательныйСрок': 3, 'Оклад': 26510, 'Организация': {'ДолжностьРуководителя': 'Индивидуальный предприниматель', 'ИНН_Орг': '773713028880', 'КПП_Орг': '', 'НаименованиеПолное': 'Индивидуальный предприниматель Тюрин Сергей Николаевич', 'НаименованиеСокращенное': 'ИП Тюрин Сергей Николаевич', 'ОГРН(ИП)': '315774600154055', 'Руководитель': 'Тюрин Сергей Николаевич', 'ФактическийАдрес': '115582, Москва г, Шипиловский проезд, дом 63, корпус 1, квартира 574', 'ЮридическийАдрес': '115582, Москва г, Шипиловский проезд, дом 63, корпус 1, квартира 574'}, 'ОснованиеПодписи': '', 'Отпуск': '28', 'Подразделение': 'Зал', 'ТкНомер': '9-тр', 'ФИО': 'Кривобок Дарья Анатольевна'}, 'personal_info': {'Адрес_проживания': 'РОССИЯ, 109382, Город Москва, вн.тер.г. муниципальный округ Люблино, ул Верхние Поля, д. 7, к. 3, кв. 77', 'Адрес_прописка': 'РОССИЯ, 109382, Город Москва, вн.тер.г. муниципальный округ Люблино, ул Верхние Поля, д. 7, к. 3, кв. 77', 'ДатаРождения': '22.06.2003', 'ИНН': '770106526206', 'МестоРождения': '0,Москва,,,', 'СНИЛС': '185-450-492 88', 'ФИО': 'Кривобок Дарья Анатольевна'}}


# def nda_sample_former(input_sample: str, output_result: str, parsed_data_1c: dict) -> str:
#     NDA_INPUT_WORDS: dict = {
#         'Полное_наименование': parsed_data_1c['job_info']['Организация']['НаименованиеПолное'],
#         '(Краткое_наименование)': parsed_data_1c['job_info']['Организация']['НаименованиеСокращенное'],
#         'Краткое_наименование': parsed_data_1c['job_info']['Организация']['НаименованиеСокращенное'],
#         'работник_Краткое_наименование': parsed_data_1c['job_info']['Организация']['НаименованиеСокращенное'],
#         'Договор': parsed_data_1c['job_info']['ТкНомер'],
#         'Дата_приказа_о_приеме': parsed_data_1c['job_info']['ДатаПриемаНаРаботу'],
#         'работодатель_Краткое_наименование': parsed_data_1c['job_info']['Организация']['НаименованиеСокращенное'],
#         'конфиденциальности_Краткое_наименование': parsed_data_1c['job_info']['Организация']['НаименованиеСокращенное']
#     }
#
#     try:
#         doc = Document(input_sample)
#
#         paragraph_count: int = len(doc.paragraphs)
#
#         for i in range(paragraph_count):
#             current_paragraph = doc.paragraphs[i]
#             current_txt = doc.paragraphs[i].text
#             # print(current_txt)
#
#             if '<Полное наименование>' in current_txt:
#                 updated_text = current_txt.replace('<Полное наименование>',
#                                                    NDA_INPUT_WORDS['Полное_наименование'])
#                 doc.paragraphs[i].text = updated_text
#                 doc.paragraphs[i].style = current_paragraph.style
#
#                 current_txt = updated_text
#
#             if current_txt == '(<Краткое наименование>)':
#                 updated_text = current_txt.replace('<Краткое наименование>',
#                                                    NDA_INPUT_WORDS['(Краткое_наименование)'])
#                 doc.paragraphs[i].text = updated_text
#                 doc.paragraphs[i].style = current_paragraph.style
#
#                 current_txt = updated_text
#
#             if '\n<Краткое наименование>' in current_txt:
#                 updated_text = current_txt.replace('<Краткое наименование>',
#                                                    NDA_INPUT_WORDS['Краткое_наименование'])
#                 doc.paragraphs[i].text = updated_text
#                 doc.paragraphs[i].style = current_paragraph.style
#
#                 current_txt = updated_text
#
#             if 'работником <Краткое наименование>' in current_txt:
#                 updated_text = current_txt.replace('работником <Краткое наименование>',
#                                                    NDA_INPUT_WORDS['работник_Краткое_наименование'])
#                 doc.paragraphs[i].text = updated_text
#                 doc.paragraphs[i].style = current_paragraph.style
#
#                 current_txt = updated_text
#
#             if '<Договор №>' in current_txt:
#                 updated_text = current_txt.replace('<Договор №>',
#                                                    NDA_INPUT_WORDS['Договор'])
#                 doc.paragraphs[i].text = updated_text
#                 doc.paragraphs[i].style = current_paragraph.style
#
#                 current_txt = updated_text
#
#             if '<Дата приказа о приеме>' in current_txt:
#                 updated_text = current_txt.replace('<Дата приказа о приеме>',
#                                                    NDA_INPUT_WORDS['Дата_приказа_о_приеме'])
#                 doc.paragraphs[i].text = updated_text
#                 doc.paragraphs[i].style = current_paragraph.style
#
#                 current_txt = updated_text
#
#             if '<Краткое наименование> (далее – «Работодатель»)' in current_txt:
#                 updated_text = current_txt.replace('<Краткое наименование>',
#                                                    NDA_INPUT_WORDS['работодатель_Краткое_наименование'])
#                 doc.paragraphs[i].text = updated_text
#                 doc.paragraphs[i].style = current_paragraph.style
#
#                 current_txt = updated_text
#
#             if '(коммерческой тайне) <Краткое наименование>' in current_txt:
#                 updated_text = current_txt.replace('<Краткое наименование>',
#                                                    NDA_INPUT_WORDS['конфиденциальности_Краткое_наименование'])
#                 doc.paragraphs[i].text = updated_text
#                 doc.paragraphs[i].style = current_paragraph.style
#
#                 current_txt = updated_text
#
#         doc.save(output_result)
#         return output_result.replace('nda/output_result/', '')
#
#     except Exception as e:
#         return f'Ошибка во время формирования файла: {e}'


def nda_sample_former(input_sample: str, output_result: str, parsed_data_1c: dict) -> str:
    nda_input_words: dict = {
        'Полное наименование': parsed_data_1c['job_info']['Организация']['НаименованиеПолное'],
        'Краткое наименование': parsed_data_1c['job_info']['Организация']['НаименованиеСокращенное'],
        'Договор №': parsed_data_1c['job_info']['ТкНомер'],
        'Дата приказа о приеме': parsed_data_1c['job_info']['ДатаПриемаНаРаботу'],
    }

    try:
        doc = Document(input_sample)
        keys = docx_get_keys(doc, pattern='<>')
        print(keys)
        docx_replace(doc, '<>', **expand_relative_month_keys(nda_input_words, keys))

        doc.save(output_result)
        return output_result

    except Exception as e:
        return f'Ошибка во время формирования файла: {e}'


if __name__ == '__main__':
    # nda_input_words: dict = {
    #     'Полное наименование': parsed_data_1c['job_info']['Организация']['НаименованиеПолное'],
    #     'Краткое наименование': parsed_data_1c['job_info']['Организация']['НаименованиеСокращенное'],
    #     'Договор №': parsed_data_1c['job_info']['ТкНомер'],
    #     'Дата приказа о приеме': parsed_data_1c['job_info']['ДатаПриемаНаРаботу'],
    # }
    #
    #
    # doc = Document("input_sample/Обязательство о неразглашении коммерческой тайны.docx")
    # keys = docx_get_keys(doc, pattern='<>')
    # docx_replace(doc, '<>', **nda_input_words)
    #
    # doc.save("output_result/output_Обязательство о неразглашении коммерческой тайны.docx")
    pass
